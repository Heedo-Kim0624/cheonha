from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Side, Border
from openpyxl.utils import get_column_letter


OUT_DIR = Path(__file__).resolve().parents[1] / "frontend" / "public" / "fleet-management" / "templates"


HEADER_FILL = PatternFill("solid", fgColor="111827")
REQUIRED_FILL = PatternFill("solid", fgColor="EAF7D1")
OPTIONAL_FILL = PatternFill("solid", fgColor="F3F6FA")
TITLE_FILL = PatternFill("solid", fgColor="F7FAEF")
THIN = Side(style="thin", color="D8DEE9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_sheet(ws, title, description, columns, examples):
    ws.title = "upload_template"
    ws.freeze_panes = "A5"

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))
    ws["A1"] = title
    ws["A2"] = description
    ws["A1"].font = Font(bold=True, size=15, color="111827")
    ws["A2"].font = Font(size=10, color="5F6B7A")
    ws["A1"].fill = TITLE_FILL
    ws["A2"].fill = TITLE_FILL
    ws["A1"].alignment = Alignment(vertical="center")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 25
    ws.row_dimensions[2].height = 35

    for col_idx, (name, required, width, note) in enumerate(columns, 1):
        cell = ws.cell(row=4, column=col_idx, value=name)
        cell.fill = REQUIRED_FILL if required else OPTIONAL_FILL
        cell.font = Font(bold=True, color="111827")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

        note_cell = ws.cell(row=5, column=col_idx, value=note)
        note_cell.fill = PatternFill("solid", fgColor="FBFCFE")
        note_cell.font = Font(size=9, color="64748B")
        note_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        note_cell.border = BORDER

    for row_idx, example in enumerate(examples, 6):
        for col_idx, value in enumerate(example, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(value, int):
                cell.number_format = "#,##0"

    for row in range(1, 6 + len(examples)):
        ws.row_dimensions[row].height = 28

    ws.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{6 + len(examples) - 1}"


def save_template(filename, title, description, columns, examples):
    wb = Workbook()
    ws = wb.active
    style_sheet(ws, title, description, columns, examples)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_DIR / filename)


def main():
    save_template(
        "accident_register_template.xlsx",
        "사고관리대장 업로드 양식",
        "사고 1건당 1행으로 작성합니다. 사고번호 기준으로 신규 등록 또는 기존 사고를 업데이트합니다.",
        [
            ("사고번호", True, 16, "예: ACC-2026-001"),
            ("차량번호", True, 16, "예: 12가3456"),
            ("실제 차량 ID", False, 18, "시스템 차량 ID"),
            ("차대번호", False, 22, "VIN"),
            ("운전자", False, 14, "운전자명"),
            ("사고일시", True, 20, "YYYY-MM-DD HH:MM"),
            ("사고장소", False, 28, "주소 또는 위치"),
            ("담보", False, 14, "대물/자차/대인"),
            ("피해자/물", False, 18, "피해 대상"),
            ("보상금액", False, 14, "숫자만 입력"),
            ("지급금액", False, 14, "숫자만 입력"),
            ("담당자", False, 14, "담당자명"),
            ("처리상태", False, 14, "진행중/종결"),
            ("사고내용", False, 36, "상세 설명"),
        ],
        [
            (
                "ACC-2026-001",
                "12가3456",
                "VEH-1001",
                "KMFXKS7BPNU000001",
                "홍길동",
                "2026-06-23 09:30",
                "서울시 강남구 테헤란로",
                "대물",
                "상대 차량",
                1250000,
                500000,
                "김담당",
                "진행중",
                "후미 추돌 사고",
            )
        ],
    )

    save_template(
        "claim_detail_template.xlsx",
        "보상상세내역 업로드 양식",
        "지급처 1곳당 1행으로 작성합니다. 사고번호와 순번을 기준으로 보상 상세내역을 관리합니다.",
        [
            ("사고번호", True, 16, "사고관리대장 사고번호"),
            ("순번", True, 10, "1, 2, 3..."),
            ("차량번호", False, 16, "예: 12가3456"),
            ("담보", False, 14, "대물/자차/대인"),
            ("피해자/물", False, 18, "피해 대상"),
            ("지급처", True, 24, "정비소/피해자/보험사"),
            ("보상금액", False, 14, "숫자만 입력"),
            ("지급금액", True, 14, "숫자만 입력"),
            ("담당자", False, 14, "담당자명"),
            ("처리상태", False, 14, "진행중/종결"),
            ("비고", False, 36, "특이사항"),
        ],
        [
            (
                "ACC-2026-001",
                1,
                "12가3456",
                "대물",
                "상대 차량",
                "한빛자동차공업사",
                1250000,
                500000,
                "김담당",
                "진행중",
                "1차 지급",
            )
        ],
    )


if __name__ == "__main__":
    main()
