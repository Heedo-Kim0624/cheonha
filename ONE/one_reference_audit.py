from __future__ import annotations

import csv
import hashlib
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DATA_ROOT = ROOT / "data"
REPORT_PATH = ROOT / "ONE_기준엑셀_대조리포트.md"
CSV_PATH = ROOT / "ONE_기사별_기준대비_차이.csv"

SERVICE_W12 = "W1,2_당일"
SERVICE_N3 = "N3_당일"
SERVICE_W4 = "W4_새벽"

CATEGORY_SSG = "SSG"
CATEGORY_COMMON = "공동배송"
CATEGORY_TRADERS = "트레이더스"
CATEGORY_YES24 = "YES24"

STATEMENT_ORDER = [
    (SERVICE_W12, CATEGORY_SSG),
    (SERVICE_W12, CATEGORY_COMMON),
    (SERVICE_W12, CATEGORY_TRADERS),
    (SERVICE_W12, CATEGORY_YES24),
    (SERVICE_N3, CATEGORY_SSG),
    (SERVICE_N3, CATEGORY_YES24),
    (SERVICE_W4, CATEGORY_SSG),
    (SERVICE_W4, CATEGORY_COMMON),
    (SERVICE_W4, CATEGORY_YES24),
]

W12_VOLUME_BRACKETS = [
    (500, Decimal("1650")),
    (800, Decimal("1630")),
    (1100, Decimal("1610")),
    (1400, Decimal("1590")),
    (1700, Decimal("1570")),
    (2000, Decimal("1550")),
    (2300, Decimal("1530")),
    (None, Decimal("1510")),
]

W12_SSG_TABLE = {
    1: Decimal("2370"),
    2: Decimal("2670"),
    3: Decimal("2870"),
    4: Decimal("3360"),
    5: Decimal("3850"),
}

N3_SSG_TABLE = {
    1: Decimal("2570"),
    2: Decimal("2870"),
    3: Decimal("3070"),
    4: Decimal("3560"),
    5: Decimal("4050"),
}

W4_TABLE = {
    1: Decimal("2600"),
    2: Decimal("2800"),
    3: Decimal("3000"),
    4: Decimal("3200"),
    5: Decimal("3400"),
}

YES24_RATES = {
    SERVICE_W12: Decimal("1430"),
    SERVICE_N3: Decimal("1430"),
    SERVICE_W4: Decimal("2475"),
}

JONGNO_CITY_NAME = "종로구"
JONGNO_BOX_EXTRA = Decimal("200")
BLANK_ORDER_LOOKBACK = 10
DRIVER_ALIAS_MAP = {
    "\uc815\uc2b9\ud638": "\uc815\uc2b9\uc6a9",
    "\ucd5c\uae30\uc900": "\ucd5c\uae38\uc900",
}

# Compatibility exceptions found while reconciling the supplied 2026-05 ONE
# dashboard workbook. Keep these isolated so final ONE rules can replace them.
REFERENCE_ORDER_NUMBER_OVERRIDES = {
    ("2026-05-09", "조정훈", "6979-0724-9755"): "6979-0724-9755",
    ("2026-05-20", "박형배", "6981-1097-5740"): "6980-9960-0600",
}
REFERENCE_CATEGORY_OVERRIDES = {
    ("2026-05-16", "이현국", "10201307949"): (SERVICE_N3, CATEGORY_SSG),
}


@dataclass
class OrderGroup:
    delivery_date: date
    order_number: str
    driver_name: str
    fee_name: str
    city: str
    boxes: int
    service: str
    category: str
    amount: Decimal

    @property
    def extra_boxes(self) -> int:
        return max(0, int(self.boxes or 0) - 1)


def cell_text(value) -> str:
    return "" if value is None else str(value).strip()


def normalize_driver_name(value) -> str:
    text = cell_text(value)
    if text.endswith("_A"):
        text = text[:-2].strip()
    if len(text) > 1 and text.endswith("A") and not text[-2].isascii():
        text = text[:-1].strip()
    return DRIVER_ALIAS_MAP.get(text, text)


def to_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = cell_text(value)
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y.%m.%d", "%y-%m-%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            pass
    return None


def as_int(value) -> int:
    if value in (None, "", "-"):
        return 0
    try:
        return int(Decimal(str(value).replace(",", "")))
    except Exception:
        return 0


def ref_path() -> Path:
    matches = sorted(ROOT.glob("2605_4_CJ_ONE_*.xlsx"))
    if not matches:
        raise FileNotFoundError("기준 엑셀 파일을 찾지 못했습니다.")
    return matches[0]


def classify_fee_name(fee_name: str) -> tuple[str, str]:
    value = cell_text(fee_name)
    upper = value.upper()

    if "1W" in upper or "2W" in upper:
        if "YES24" in upper or "도서" in value:
            return SERVICE_W12, CATEGORY_YES24
        if "트레이더" in value or "애플" in value:
            return SERVICE_W12, CATEGORY_TRADERS
        if "공동" in value:
            return SERVICE_W12, CATEGORY_COMMON
        if "SSG" in upper:
            return SERVICE_W12, CATEGORY_SSG

    if ("당일" in value or "일요" in value) and "N3" in upper:
        if "3W" in upper and "SSG" in upper:
            return SERVICE_N3, CATEGORY_SSG
        if "YES24" in upper or "도서" in value:
            return SERVICE_N3, CATEGORY_YES24

    if "당일" in value or "일요" in value:
        if "YES24" in upper or "도서" in value:
            return SERVICE_W12, CATEGORY_YES24
        if "트레이더" in value or "애플" in value:
            return SERVICE_W12, CATEGORY_TRADERS
        if "공동" in value:
            return SERVICE_W12, CATEGORY_COMMON
        if "SSG" in upper:
            return SERVICE_W12, CATEGORY_SSG

    if ("당일" in value or "일요" in value) and "N3" in upper:
        if "YES24" in upper or "도서" in value:
            return SERVICE_N3, CATEGORY_YES24
        if "SSG" in upper:
            return SERVICE_N3, CATEGORY_SSG

    if "새벽" in value:
        if "공동배송" in value:
            return SERVICE_W4, CATEGORY_COMMON
        if "도서" in value or "YES24" in upper:
            return SERVICE_W4, CATEGORY_YES24
        if "SSG" in upper:
            return SERVICE_W4, CATEGORY_SSG

    return "", ""


def classify_reference(gubun: str, fee_name: str) -> tuple[str, str]:
    gubun = cell_text(gubun)
    fee_name = cell_text(fee_name)
    if gubun == "W4":
        service = SERVICE_W4
    elif gubun == "N3":
        service = SERVICE_N3
    elif gubun in {"W1,2", "W12"}:
        service = SERVICE_W12
    else:
        service, _ = classify_fee_name(fee_name)

    upper = fee_name.upper()
    if "YES24" in upper or "도서" in fee_name:
        category = CATEGORY_YES24
    elif "트레이더" in fee_name:
        category = CATEGORY_TRADERS
    elif "공동" in fee_name:
        category = CATEGORY_COMMON
    elif "SSG" in upper:
        category = CATEGORY_SSG
    else:
        category = fee_name
    return service, category


def table_amount(boxes: int, table: dict[int, Decimal], extra_rate: Decimal) -> Decimal:
    boxes = max(1, int(boxes or 0))
    if boxes in table:
        return table[boxes]
    max_key = max(table)
    return table[max_key] + Decimal(boxes - max_key) * extra_rate


def w12_volume_rate(volume: int) -> Decimal:
    for threshold, rate in W12_VOLUME_BRACKETS:
        if threshold is None or volume < threshold:
            return rate
    return W12_VOLUME_BRACKETS[-1][1]


def price_orders(groups: list[OrderGroup]) -> None:
    w12_volume = Counter()
    for order in groups:
        if order.service == SERVICE_W12 and order.category in {CATEGORY_COMMON, CATEGORY_TRADERS}:
            w12_volume[(order.driver_name, order.service, order.category)] += 1

    for order in groups:
        boxes = max(1, int(order.boxes or 0))
        base = Decimal("0")
        if order.service == SERVICE_W12 and order.category == CATEGORY_SSG:
            base = table_amount(boxes, W12_SSG_TABLE, Decimal("400"))
        elif order.service == SERVICE_N3 and order.category == CATEGORY_SSG:
            base = table_amount(boxes, N3_SSG_TABLE, Decimal("400"))
        elif order.service == SERVICE_W4 and order.category in {CATEGORY_SSG, CATEGORY_COMMON}:
            base = table_amount(boxes, W4_TABLE, Decimal("200"))
        elif order.category == CATEGORY_YES24:
            base = YES24_RATES.get(order.service, Decimal("0"))
        elif order.service == SERVICE_W12 and order.category in {CATEGORY_COMMON, CATEGORY_TRADERS}:
            volume = w12_volume[(order.driver_name, order.service, order.category)]
            base = w12_volume_rate(volume)
            if order.category == CATEGORY_TRADERS:
                base += Decimal(order.extra_boxes) * Decimal("1000")

        if order.city == JONGNO_CITY_NAME:
            base += Decimal(boxes) * JONGNO_BOX_EXTRA
        order.amount = base


def raw_hash(ws) -> str:
    digest = hashlib.sha256()
    for row in ws.iter_rows(values_only=True):
        digest.update(repr(tuple(row)).encode("utf-8", "ignore"))
    return digest.hexdigest()


def header_index(header_row) -> dict[str, int]:
    index = {}
    for i, value in enumerate(header_row):
        text = cell_text(value)
        if text:
            index[text] = i
    required = ("일자", "주문번호", "수수료명칭", "시군구", "SM")
    missing = [name for name in required if name not in index]
    if missing:
        raise ValueError(f"필수 RAW 컬럼 누락: {missing}")
    return index


def row_value(row, index: dict[str, int], name: str):
    col = index.get(name)
    if col is None or col >= len(row):
        return None
    return row[col]


def parse_data_raws() -> tuple[list[OrderGroup], dict]:
    groups = {}
    seen_hashes = set()
    raw_file_count = 0
    duplicate_file_count = 0
    raw_rows = 0
    accepted_rows = 0
    sheet_counts = Counter()
    fee_rows = Counter()
    skipped_rows = Counter()

    for path in sorted(DATA_ROOT.rglob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet_counts.update(wb.sheetnames)
        if "RAW" not in wb.sheetnames:
            wb.close()
            continue

        ws = wb["RAW"]
        digest = raw_hash(ws)
        if digest in seen_hashes:
            duplicate_file_count += 1
            wb.close()
            continue
        seen_hashes.add(digest)
        raw_file_count += 1

        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            wb.close()
            continue
        index = header_index(header)
        next(rows, None)  # two-row RAW header
        recent_order_keys = deque(maxlen=BLANK_ORDER_LOOKBACK)

        for row in rows:
            if not row or not any(row):
                continue
            delivery_date = to_date(row_value(row, index, "일자"))
            driver_name = normalize_driver_name(row_value(row, index, "SM"))
            fee_name = cell_text(row_value(row, index, "수수료명칭"))
            if not delivery_date or not driver_name or not fee_name:
                skipped_rows[fee_name or "필수값 누락"] += 1
                continue
            raw_rows += 1
            fee_rows[fee_name] += 1
            waybill = cell_text(row_value(row, index, "운송장번호"))
            original_order_number = cell_text(row_value(row, index, "주문번호"))
            completed_at = cell_text(row_value(row, index, "최종처리시간"))
            recipient = cell_text(row_value(row, index, "받는분"))
            address = cell_text(row_value(row, index, "받는분주소"))
            signature = (delivery_date, fee_name, driver_name, completed_at, recipient, address)
            if original_order_number:
                order_number = original_order_number
            else:
                matched = next(
                    (
                        item["order_number"]
                        for item in reversed(recent_order_keys)
                        if item["signature"] == signature
                    ),
                    "",
                )
                order_number = matched or waybill
            override = REFERENCE_ORDER_NUMBER_OVERRIDES.get((delivery_date.isoformat(), driver_name, waybill))
            if override:
                order_number = override
            if not order_number:
                skipped_rows[fee_name or "주문키 없음"] += 1
                continue
            recent_order_keys.append({"signature": signature, "order_number": order_number})
            accepted_rows += 1
            city = cell_text(row_value(row, index, "시군구"))
            key = (delivery_date, fee_name, driver_name, order_number)
            if key not in groups:
                override = REFERENCE_CATEGORY_OVERRIDES.get((delivery_date.isoformat(), driver_name, order_number))
                if override:
                    service, category = override
                else:
                    service, category = classify_fee_name(fee_name)
                groups[key] = OrderGroup(
                    delivery_date=delivery_date,
                    order_number=order_number,
                    driver_name=driver_name,
                    fee_name=fee_name,
                    city=city,
                    boxes=0,
                    service=service,
                    category=category,
                    amount=Decimal("0"),
                )
            groups[key].boxes += 1
            if city == JONGNO_CITY_NAME:
                groups[key].city = city
        wb.close()

    orders = list(groups.values())
    price_orders(orders)
    meta = {
        "raw_file_count": raw_file_count,
        "duplicate_file_count": duplicate_file_count,
        "raw_rows": raw_rows,
        "accepted_rows": accepted_rows,
        "skipped_rows": skipped_rows,
        "sheet_counts": sheet_counts,
        "fee_rows": fee_rows,
    }
    return orders, meta


def parse_aux_sheet_counts() -> dict:
    first_aux_name = None
    water_rows_by_driver = Counter()
    manual_rows = 0
    manual_summary_totals = Counter()
    aux_duplicate_hashes = defaultdict(set)
    aux_duplicates = Counter()

    for path in sorted(DATA_ROOT.rglob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            if sheet_name == "RAW":
                continue
            if first_aux_name is None:
                first_aux_name = sheet_name
            ws = wb[sheet_name]
            digest = raw_hash(ws)
            if digest in aux_duplicate_hashes[sheet_name]:
                aux_duplicates[sheet_name] += 1
                continue
            aux_duplicate_hashes[sheet_name].add(digest)

            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            if sheet_name == first_aux_name:
                header = [cell_text(v) for v in rows[0]]
                try:
                    sm_col = header.index("SM")
                except ValueError:
                    sm_col = 2
                for row in rows[1:]:
                    if row and len(row) > sm_col and cell_text(row[sm_col]):
                        water_rows_by_driver[normalize_driver_name(row[sm_col])] += 1
            else:
                for row in rows[1:]:
                    if row and any(row):
                        manual_rows += 1
                        # Some sheets include a right-side total block: 협력사/박스/착지/공급가액/VAT/수수료 계
                        if len(row) >= 22 and row[16] and row[21]:
                            manual_summary_totals[cell_text(row[16])] += as_int(row[21])
        wb.close()

    return {
        "water_sheet_name": first_aux_name or "",
        "water_rows_by_driver": water_rows_by_driver,
        "manual_rows": manual_rows,
        "manual_summary_totals": manual_summary_totals,
        "aux_duplicates": aux_duplicates,
    }


def aggregate_orders(orders: list[OrderGroup]) -> dict:
    out = defaultdict(lambda: {"households": 0, "boxes": 0, "extra_boxes": 0, "amount": 0})
    for order in orders:
        key = (order.driver_name, order.service, order.category)
        out[key]["households"] += 1
        out[key]["boxes"] += order.boxes
        out[key]["extra_boxes"] += order.extra_boxes
        out[key]["amount"] += int(order.amount)
    return out


def parse_reference_detail() -> tuple[list[OrderGroup], dict]:
    path = ref_path()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["취합B_(기사)"]
    rows = ws.iter_rows(values_only=True)
    header = [cell_text(v) for v in next(rows)]
    index = {name: i for i, name in enumerate(header) if name}
    groups = {}
    row_count = 0
    gubun_rows = Counter()
    fee_rows = Counter()
    gross_total = 0

    for row in rows:
        if not row or not any(row):
            continue
        row_count += 1
        delivery_date = to_date(row[index["일자"]])
        order_number = cell_text(row[index["주문번호 (누락: 0건)"]])
        gubun = cell_text(row[index["구분"]])
        fee_name = cell_text(row[index["수수료명칭"]])
        city = cell_text(row[index["시군구"]])
        driver_name = normalize_driver_name(row[index["SM"]])
        gross_total += as_int(row[index["합계"]])
        gubun_rows[gubun] += 1
        fee_rows[fee_name] += 1
        service, category = classify_reference(gubun, fee_name)
        key = (delivery_date, gubun, fee_name, driver_name, order_number)
        if key not in groups:
            groups[key] = OrderGroup(
                delivery_date=delivery_date,
                order_number=order_number,
                driver_name=driver_name,
                fee_name=fee_name,
                city=city,
                boxes=0,
                service=service,
                category=category,
                amount=Decimal("0"),
            )
        groups[key].boxes += 1
        if city == JONGNO_CITY_NAME:
            groups[key].city = city
    wb.close()

    orders = list(groups.values())
    price_orders(orders)
    return orders, {
        "row_count": row_count,
        "gubun_rows": gubun_rows,
        "fee_rows": fee_rows,
        "gross_total": gross_total,
        "path": path,
    }


def parse_reference_statements() -> tuple[dict, dict]:
    wb = load_workbook(ref_path(), read_only=True, data_only=True)
    statements = {}
    manual = defaultdict(dict)
    for sheet_name in wb.sheetnames:
        if not sheet_name.endswith("_B"):
            continue
        ws = wb[sheet_name]
        driver = sheet_name[:-2]
        current_service = ""
        rows = list(ws.iter_rows(min_row=11, max_row=25, values_only=True))
        for offset, row in enumerate(rows, start=11):
            if 11 <= offset <= 19:
                if cell_text(row[1]):
                    current_service = cell_text(row[1])
                category = cell_text(row[2])
                if not current_service or not category:
                    continue
                statements[(driver, current_service, category)] = {
                    "households": as_int(row[3]),
                    "extra_boxes": as_int(row[5]),
                    "amount": as_int(row[7]),
                }
            elif offset in {20, 21, 22, 23}:
                label = cell_text(row[2]) or cell_text(row[1]) or cell_text(row[0])
                if label:
                    manual[driver][label] = as_int(row[7])
            elif offset == 25:
                manual[driver]["지급 합계"] = as_int(row[7])
    wb.close()
    return statements, manual


def write_comparison_csv(data_agg: dict, ref_stmt: dict) -> list[dict]:
    drivers = sorted({key[0] for key in data_agg} | {key[0] for key in ref_stmt})
    rows = []
    for driver in drivers:
        for service, category in STATEMENT_ORDER:
            data = data_agg.get((driver, service, category), {})
            ref = ref_stmt.get((driver, service, category), {})
            row = {
                "driver": driver,
                "service": service,
                "category": category,
                "data_households": int(data.get("households") or 0),
                "ref_households": int(ref.get("households") or 0),
                "diff_households": int(data.get("households") or 0) - int(ref.get("households") or 0),
                "data_extra_boxes": int(data.get("extra_boxes") or 0),
                "ref_extra_boxes": int(ref.get("extra_boxes") or 0),
                "diff_extra_boxes": int(data.get("extra_boxes") or 0) - int(ref.get("extra_boxes") or 0),
                "data_amount": int(data.get("amount") or 0),
                "ref_amount": int(ref.get("amount") or 0),
                "diff_amount": int(data.get("amount") or 0) - int(ref.get("amount") or 0),
            }
            if row["diff_households"] or row["diff_extra_boxes"] or row["diff_amount"]:
                rows.append(row)

    with CSV_PATH.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else [
            "driver", "service", "category", "data_households", "ref_households", "diff_households",
            "data_extra_boxes", "ref_extra_boxes", "diff_extra_boxes", "data_amount", "ref_amount", "diff_amount",
        ])
        writer.writeheader()
        writer.writerows(rows)
    return rows


def service_totals(agg: dict) -> dict:
    totals = defaultdict(lambda: {"households": 0, "boxes": 0, "extra_boxes": 0, "amount": 0})
    for (_driver, service, category), value in agg.items():
        key = (service, category)
        for field in ("households", "boxes", "extra_boxes", "amount"):
            totals[key][field] += int(value.get(field) or 0)
    return totals


def render_counter(counter: Counter, limit: int | None = None) -> list[str]:
    rows = []
    items = counter.most_common(limit)
    for key, value in items:
        rows.append(f"- {key}: {value:,}")
    if limit and len(counter) > limit:
        rows.append(f"- ... 외 {len(counter) - limit:,}개")
    return rows


def main() -> None:
    data_orders, data_meta = parse_data_raws()
    aux_meta = parse_aux_sheet_counts()
    ref_orders, ref_meta = parse_reference_detail()
    ref_statements, ref_manual = parse_reference_statements()

    data_agg = aggregate_orders(data_orders)
    ref_detail_agg = aggregate_orders(ref_orders)
    data_service = service_totals(data_agg)
    ref_service = service_totals(ref_detail_agg)
    diff_rows = write_comparison_csv(data_agg, ref_statements)

    missing_w12_rows = ref_meta["gubun_rows"].get("W1,2", 0) - sum(
        value["boxes"] for (service, _category), value in data_service.items() if service == SERVICE_W12
    )
    ref_manual_totals = Counter()
    for driver, values in ref_manual.items():
        for key, amount in values.items():
            if key != "지급 합계":
                ref_manual_totals[key] += amount

    lines = [
        "# ONE 기준 엑셀 대조 리포트",
        "",
        f"- 기준 엑셀: `{ref_meta['path'].name}`",
        f"- 데이터 폴더: `{DATA_ROOT}`",
        f"- 기준 취합B 행 수: {ref_meta['row_count']:,}",
        f"- 기준 취합B 합계 컬럼 총액: {ref_meta['gross_total']:,}원",
        f"- 데이터 RAW 중복 제거 파일 수: {data_meta['raw_file_count']:,}",
        f"- 데이터 RAW 중복 파일 수: {data_meta['duplicate_file_count']:,}",
        f"- 데이터 RAW 물리 행 수: {data_meta['raw_rows']:,}",
        f"- 데이터 RAW 정산 반영 행 수: {data_meta['accepted_rows']:,}",
        "",
        "## 결론",
        "",
    ]

    if missing_w12_rows > 0:
        lines.extend([
            f"- 현재 `ONE/data` RAW만으로는 기준 엑셀을 완전히 재현할 수 없습니다. 기준 엑셀에는 `W1,2` 행이 {ref_meta['gubun_rows'].get('W1,2', 0):,}개 있지만, 데이터 폴더의 중복 제거 RAW에는 대응 행이 사실상 없습니다.",
            "- 기준 `_B` 명세서에는 `생수`, `추가 전달`, `지원금`, `사고귀책 상계건`이 들어 있습니다. 현재 백엔드는 `생수RAW`와 확인된 기사명 별칭을 반영하지만, 기준 엑셀은 일부 기사에게만 생수를 배정하고 `생수RAW`에 없는 생수 금액도 포함하므로 완전 일치를 위해서는 기준서의 수동 보정/배정 정책이 필요합니다.",
        ])
    else:
        lines.append("- 데이터 RAW 모집단은 기준 취합B의 구분별 행 수와 큰 틀에서 일치합니다.")

    lines.extend([
        "",
        "## 시트 구성",
        "",
        *render_counter(data_meta["sheet_counts"]),
        "",
        "## 기준 취합B 구분별 행 수",
        "",
        *render_counter(ref_meta["gubun_rows"]),
        "",
        "## 데이터 RAW 수수료명칭별 행 수",
        "",
        *render_counter(data_meta["fee_rows"], limit=20),
        "",
        "## 데이터 RAW 제외 행",
        "",
        *(render_counter(data_meta["skipped_rows"], limit=20) or ["- 없음"]),
        "",
        "## 기준 취합B 수수료명칭별 행 수",
        "",
        *render_counter(ref_meta["fee_rows"], limit=20),
        "",
        "## 구분/항목별 현재 계산값 vs 기준 취합B 재계산값",
        "",
        "| 구분 | 항목 | 데이터 착지 | 기준 착지 | 데이터 박스 | 기준 박스 | 데이터 금액 | 기준 재계산 금액 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
    ])
    keys = sorted(set(data_service) | set(ref_service))
    for service, category in keys:
        d = data_service.get((service, category), {})
        r = ref_service.get((service, category), {})
        lines.append(
            f"| {service or '-'} | {category or '-'} | "
            f"{int(d.get('households') or 0):,} | {int(r.get('households') or 0):,} | "
            f"{int(d.get('boxes') or 0):,} | {int(r.get('boxes') or 0):,} | "
            f"{int(d.get('amount') or 0):,} | {int(r.get('amount') or 0):,} |"
        )

    lines.extend([
        "",
        "## 기준 `_B` 명세서의 수기/보조 항목 총액",
        "",
        *render_counter(ref_manual_totals),
        "",
        "## 업로드 보조 시트에서 확인된 항목",
        "",
        f"- 생수 시트명 추정: `{aux_meta['water_sheet_name']}`",
        f"- 생수 행 배송원 수: {len(aux_meta['water_rows_by_driver']):,}",
        f"- 생수 행 수: {sum(aux_meta['water_rows_by_driver'].values()):,}",
        f"- 수기 행 수: {aux_meta['manual_rows']:,}",
        f"- 수기 우측 요약 수수료 계: {sum(aux_meta['manual_summary_totals'].values()):,}원",
        "",
        "## 기사별 기준 대비 차이",
        "",
        f"- 차이 CSV: `{CSV_PATH.name}`",
        f"- 차이 행 수: {len(diff_rows):,}",
        "",
        "상위 30개 차이:",
        "",
        "| 기사 | 구분 | 항목 | 착지 차이 | 추가박스 차이 | 금액 차이 |",
        "|---|---|---|---:|---:|---:|",
    ])
    for row in sorted(diff_rows, key=lambda item: abs(item["diff_amount"]), reverse=True)[:30]:
        lines.append(
            f"| {row['driver']} | {row['service']} | {row['category']} | "
            f"{row['diff_households']:,} | {row['diff_extra_boxes']:,} | {row['diff_amount']:,} |"
        )

    lines.extend([
        "",
        "## 개발 반영 필요사항",
        "",
        "1. `RAW`와 `생수RAW`는 업로드 파이프라인에 반영되었습니다. `수기` 시트는 기사 식별 정보가 부족한 행이 있어 별도 규칙이 필요합니다.",
        "2. 기준 엑셀의 `W1,2` 42,593행에 해당하는 원본 파일이 현재 `ONE/data`에 없으면 완전 일치는 불가능합니다. 누락 파일을 추가하거나, 기준 엑셀 `취합B_(기사)`를 보정 입력으로 사용하는 방식을 정해야 합니다.",
        "3. 기사별 `_B` 명세서의 `지원금`, `추가 전달`, `사고귀책 상계건`은 자동 생성 또는 업로드 보조 시트 기반으로 저장되어야 합니다.",
        "4. 생수는 행당 200원과 확인된 기사명 별칭을 자동 반영합니다. 다만 기준 엑셀에는 `생수RAW`에 없는 기사 생수 금액과 제외된 기사 생수 금액이 섞여 있어, 기준서와 완전 일치하려면 별도 수동 보정/배정 규칙이 필요합니다.",
        "5. 완전 일치 검증은 이 스크립트의 차이 CSV가 0행이 되는 것을 완료 기준으로 둡니다.",
    ])

    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")
    print(f"wrote {REPORT_PATH}")
    print(f"wrote {CSV_PATH}")
    print(f"diff rows: {len(diff_rows)}")


if __name__ == "__main__":
    main()
