"""
backend/apps/tracking/management/commands/seed_tracking.py

example_data/ 의 실제 1인분 배송 기록을 읽어 **한 명**의 추적 세션으로 seed.

파이프라인:
1. xlsx 로드  (GPS / BLE / 카메라)
2. GPS → Kalman+RTS 스무딩  (apps.tracking.algorithms.kalman)
3. BLE → HMM 2-state(In Vehicle / Cycle) + 사이클 내 SR→DL boundary
   (apps.tracking.algorithms.ble_state)
4. GPS 타임스탬프에 BLE state 보간하여 LocationPoint 저장
5. 사이클별 집계

사용: python manage.py seed_tracking
옵션:
  --date YYYY-MM-DD   (기본: 2026-04-17)
  --team <code>       (기본: 첫 번째 팀)
  --crew <code>       (기본: 첫 번째 crew)
  --clear             (기존 tracking 전부 삭제)
  --raw               (Kalman/BLE HMM 끄고 speed heuristic 으로 fallback)

camera_end_time 은 촬영 시각(capture)으로 간주한다.
"""
import math
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from django.core.management.base import BaseCommand
from django.db import transaction
from django.utils import timezone

from apps.accounts.models import Team
from apps.crew.models import CrewMember
from apps.tracking.models import (
    TrackingSession, LocationPoint, BleLog, CameraCapture, Cycle,
    STATE_IN_VEHICLE, STATE_SEARCHING, STATE_DELIVERING,
)
from apps.tracking.algorithms.kalman import smooth_gps_track
from apps.tracking.algorithms.ble_state import classify_three_states


DEFAULT_CREW = {'code': 'TRK001', 'name': '박철수', 'vehicle': '서울90바8677'}


def _resolve_example_dir():
    here = Path(__file__).resolve()
    candidates = [
        here.parents[5] / 'example_data',
        here.parents[4] / 'example_data',
        Path('/app/example_data'),
        Path('/example_data'),
    ]
    for c in candidates:
        if c.exists() and (c / 'driver_locations_matching_779_2026-04-17.xlsx').exists():
            return c
    return candidates[0]


EXAMPLE_DIR = _resolve_example_dir()


def haversine_m(lat1, lon1, lat2, lon2):
    R = 6371000.0
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(min(1.0, math.sqrt(a)))


def _to_datetime(v):
    if isinstance(v, datetime):
        return v
    if v in (None, ''):
        return None
    return datetime.strptime(str(v).strip(), '%Y-%m-%d %H:%M:%S')


def _to_float(v, default=0.0):
    try:
        if v in (None, ''):
            return default
        return float(v)
    except Exception:
        return default


def _to_int(v, default=0):
    try:
        if v in (None, ''):
            return default
        return int(float(v))
    except Exception:
        return default


class Command(BaseCommand):
    help = 'example_data xlsx → Kalman · BLE HMM 파이프라인으로 배송 추적 세션 seed'

    def add_arguments(self, parser):
        parser.add_argument('--date', default='2026-04-17')
        parser.add_argument('--team', default=None)
        parser.add_argument('--crew', default=None)
        parser.add_argument('--clear', action='store_true')
        parser.add_argument('--raw', action='store_true',
                            help='알고리즘 비활성화 후 단순 speed heuristic 사용')

    def handle(self, *args, **opts):
        target_date = datetime.strptime(opts['date'], '%Y-%m-%d').date()

        if opts['clear']:
            deleted, _ = TrackingSession.objects.all().delete()
            self.stdout.write(self.style.WARNING(f'기존 tracking 삭제: {deleted}건'))

        team = self._pick_team(opts.get('team'))
        self.stdout.write(f'대상 팀: {team.code} ({team.name})')

        base = self._load_base()
        self.stdout.write(
            f'기준 데이터 — GPS {len(base["gps"])} · BLE {len(base["ble"])} · CAM {len(base["cam"])}'
        )

        crew = self._ensure_crew(team, opts.get('crew'))

        with transaction.atomic():
            self._seed_for_crew(crew, base, target_date, use_algo=not opts.get('raw'))

        self.stdout.write(self.style.SUCCESS(
            f'✅ {target_date} / {team.name} / {crew.name}({crew.code}) seed 완료'
        ))

    # ------------------------------------------------------------------
    def _pick_team(self, code):
        if code:
            t = Team.objects.filter(code=code).first()
            if t:
                return t
        return Team.objects.first() or Team.objects.create(code='YDP2', name='영등포 2조')

    def _ensure_crew(self, team, crew_code=None):
        if crew_code:
            c = CrewMember.objects.filter(team=team, code=crew_code).first()
            if c:
                return c
        c = CrewMember.objects.filter(team=team, is_active=True).order_by('code').first()
        if c:
            return c
        c, _ = CrewMember.objects.get_or_create(
            code=DEFAULT_CREW['code'], team=team,
            defaults={
                'name': DEFAULT_CREW['name'],
                'vehicle_number': DEFAULT_CREW['vehicle'],
                'is_active': True,
            },
        )
        return c

    # ------------------------------------------------------------------
    def _load_base(self):
        gps_p = EXAMPLE_DIR / 'driver_locations_matching_779_2026-04-17.xlsx'
        ble_p = EXAMPLE_DIR / 'ble_logs_matching_779_2026-04-17.xlsx'
        cam_p = EXAMPLE_DIR / 'camera_logs_matching_779_2026-04-17.xlsx'

        gps = []
        wb = load_workbook(gps_p, read_only=True, data_only=True)
        for i, row in enumerate(wb.active.iter_rows(values_only=True)):
            if i == 0:
                continue
            lat = _to_float(row[3], None)
            lon = _to_float(row[4], None)
            t = _to_datetime(row[7])
            if lat is None or lon is None or t is None:
                continue
            gps.append({'time': t, 'lat': lat, 'lon': lon,
                        'acc': _to_float(row[5], 10.0),
                        'spd': _to_float(row[6], 0.0)})
        wb.close()
        gps.sort(key=lambda x: x['time'])

        ble = []
        wb = load_workbook(ble_p, read_only=True, data_only=True)
        for i, row in enumerate(wb.active.iter_rows(values_only=True)):
            if i == 0:
                continue
            t = _to_datetime(row[4])
            if t is None:
                continue
            typ = str(row[5] or '').lower()
            rssi = _to_int(row[7], None) if row[7] not in (None, '') else None
            ble.append({'time': t, 'type': typ, 'rssi': rssi})
        wb.close()
        ble.sort(key=lambda x: x['time'])

        cam = []
        wb = load_workbook(cam_p, read_only=True, data_only=True)
        for i, row in enumerate(wb.active.iter_rows(values_only=True)):
            if i == 0:
                continue
            start = _to_datetime(row[4])
            end = _to_datetime(row[5]) or start
            if end is None:
                continue
            cam.append({'started_at': start,
                        'captured_at': end,
                        'duration_ms': _to_int(row[6], 0),
                        'device': str(row[8] or '')[:64]})
        wb.close()
        cam.sort(key=lambda x: x['captured_at'])

        return {'gps': gps, 'ble': ble, 'cam': cam}

    # ------------------------------------------------------------------
    def _seed_for_crew(self, crew, base, target_date, use_algo=True):
        if not base['gps']:
            return

        base_day = base['gps'][0]['time'].date()
        shift = target_date - base_day

        def sh(t):  # naive datetime 에 shift 적용
            return t + shift

        # ---------- 카메라 시각 KST 보정 (+9h) ----------
        # example_data 의 camera 는 UTC 로 기록되어 있고 GPS 는 KST 로 기록돼 있어
        # 그대로 두면 두 시계열이 9시간 차이나서 nearest-neighbor 가 모두 첫 GPS 포인트로 snap 됨.
        kst_delta = timedelta(hours=9)
        for c in base['cam']:
            if c.get('started_at'):
                c['started_at'] = c['started_at'] + kst_delta
            c['captured_at'] = c['captured_at'] + kst_delta
        if base['cam']:
            self.stdout.write(
                f'  · camera 시각 +9h (UTC → KST) 적용: '
                f'{base["cam"][0]["captured_at"]} ~ {base["cam"][-1]["captured_at"]}'
            )

        # ---------- GPS → Kalman+RTS ----------
        gps_df = pd.DataFrame([{
            'time': sh(p['time']),
            'lat': p['lat'], 'lon': p['lon'],
            'acc': p['acc'], 'spd': p['spd'],
        } for p in base['gps']])

        t_sec = (pd.to_datetime(gps_df['time']).astype('int64') // 10 ** 9).values

        if use_algo:
            filt_lat, filt_lon = smooth_gps_track(
                gps_df['lat'].values,
                gps_df['lon'].values,
                t_sec,
                accuracy_m=gps_df['acc'].values,
                speed_kmh=gps_df['spd'].values,
            )
            # NaN 발생 포인트는 원본으로 폴백
            mask_nan = ~np.isfinite(filt_lat) | ~np.isfinite(filt_lon)
            if mask_nan.any():
                filt_lat[mask_nan] = gps_df['lat'].values[mask_nan]
                filt_lon[mask_nan] = gps_df['lon'].values[mask_nan]
            gps_df['lat_f'] = filt_lat
            gps_df['lon_f'] = filt_lon
        else:
            gps_df['lat_f'] = gps_df['lat']
            gps_df['lon_f'] = gps_df['lon']

        # ---------- BLE → 3-state ----------
        ble_df = pd.DataFrame([{
            'time': sh(b['time']),
            'rssi': b['rssi'],
            'type': b['type'],
        } for b in base['ble'] if b['type'] == 'rssi' and b['rssi'] is not None])

        state_series = None
        if use_algo and not ble_df.empty:
            state_series = classify_three_states(ble_df[['time', 'rssi']])

        # ---------- GPS 각 포인트에 state 매핑 ----------
        def state_at(ts):
            if state_series is None or state_series.empty:
                return (STATE_IN_VEHICLE, 0)
            idx = state_series.index.get_indexer([ts], method='nearest')[0]
            row = state_series.iloc[idx]
            state = {'IV': STATE_IN_VEHICLE, 'SR': STATE_SEARCHING, 'DL': STATE_DELIVERING}.get(
                row['state'], STATE_IN_VEHICLE)
            return (state, int(row['cycle_id']))

        cam_times_aw = [timezone.make_aware(sh(c['captured_at'])) for c in base['cam']]

        # 시간 index 를 aware 로
        gps_df['time_aware'] = gps_df['time'].apply(timezone.make_aware)

        # ---------- 세션 생성 ----------
        TrackingSession.objects.filter(crew_member=crew, session_date=target_date).delete()
        session = TrackingSession.objects.create(
            crew_member=crew,
            session_date=target_date,
            started_at=gps_df['time_aware'].iloc[0],
            ended_at=gps_df['time_aware'].iloc[-1],
            device_id='DEMO-DEVICE',
        )

        # LocationPoint
        points_to_create = []
        distance_m = 0.0
        iv_sec = sr_sec = dl_sec = 0
        prev_lat = prev_lon = None
        last_time = None
        fallback_near_cam = state_series is None

        for _, row in gps_df.iterrows():
            t_aw = row['time_aware']
            if fallback_near_cam:
                # raw 모드: speed + camera proximity 휴리스틱
                near = any(abs((t_aw - ct).total_seconds()) <= 60 for ct in cam_times_aw)
                if row['spd'] > 10:
                    state, cid = STATE_IN_VEHICLE, 0
                elif near:
                    state, cid = STATE_DELIVERING, 0
                else:
                    state, cid = STATE_SEARCHING, 0
            else:
                state, cid = state_at(pd.Timestamp(row['time']))

            points_to_create.append(LocationPoint(
                session=session,
                recorded_at=t_aw,
                lat=float(row['lat_f']),
                lon=float(row['lon_f']),
                accuracy_m=float(row['acc']),
                speed_kmh=float(row['spd']),
                state=state,
                cycle_id=cid or None,
            ))

            if last_time is not None:
                dt = max(0, int((t_aw - last_time).total_seconds()))
                if state == STATE_IN_VEHICLE:
                    iv_sec += dt
                elif state == STATE_SEARCHING:
                    sr_sec += dt
                else:
                    dl_sec += dt
            last_time = t_aw

            if prev_lat is not None:
                distance_m += haversine_m(prev_lat, prev_lon, row['lat_f'], row['lon_f'])
            prev_lat, prev_lon = row['lat_f'], row['lon_f']

        LocationPoint.objects.bulk_create(points_to_create, batch_size=1000)

        # ---------- Cycle 객체 ----------
        self._build_cycles(session, points_to_create)

        # ---------- BLE Log 저장 (샘플) ----------
        ble_rssi = [b for b in base['ble']
                    if b['type'] == 'rssi' and b['rssi'] is not None][:2000]
        BleLog.objects.bulk_create([
            BleLog(session=session,
                   recorded_at=timezone.make_aware(sh(b['time'])),
                   rssi=b['rssi'], status='')
            for b in ble_rssi
        ], batch_size=500)

        # ---------- Camera ----------
        # started_at 과 captured_at(=end) 사이를 4:1 로 내분한 시점의 GPS 위치를 저장
        capture_rows = []
        pts_sorted = points_to_create  # 이미 시간순
        for c in base['cam']:
            end_at = timezone.make_aware(sh(c['captured_at']))
            start_at = timezone.make_aware(sh(c['started_at'])) if c.get('started_at') else end_at
            # 내분비 4:1 (start 에서 end 쪽으로 80%)
            dt = (end_at - start_at).total_seconds()
            anchor = start_at + timedelta(seconds=dt * (4.0 / 5.0))
            nearest = min(pts_sorted,
                          key=lambda p: abs((p.recorded_at - anchor).total_seconds()),
                          default=None)
            capture_rows.append(CameraCapture(
                session=session,
                started_at=start_at,
                captured_at=end_at,
                duration_ms=c['duration_ms'],
                device=c['device'],
                lat=nearest.lat if nearest else None,
                lon=nearest.lon if nearest else None,
                cycle_id=nearest.cycle_id if nearest else None,
            ))
        CameraCapture.objects.bulk_create(capture_rows, batch_size=200)

        # 사이클별 capture_count 업데이트
        for cyc in session.cycles.all():
            n_cap = session.captures.filter(cycle_id=cyc.id).count()
            if n_cap:
                cyc.capture_count = n_cap
                cyc.save(update_fields=['capture_count'])

        # ---------- 집계 ----------
        lats = [p.lat for p in points_to_create]
        lons = [p.lon for p in points_to_create]
        session.total_seconds = iv_sec + sr_sec + dl_sec
        session.iv_seconds = iv_sec
        session.sr_seconds = sr_sec
        session.dl_seconds = dl_sec
        session.distance_m = distance_m
        session.cycle_count = session.cycles.count()
        session.bbox_min_lat = min(lats)
        session.bbox_max_lat = max(lats)
        session.bbox_min_lon = min(lons)
        session.bbox_max_lon = max(lons)
        session.save()

        self.stdout.write(
            f'  · {crew.code} {crew.name}  pts={len(points_to_create)} '
            f'cycles={session.cycle_count}  IV={iv_sec}s SR={sr_sec}s DL={dl_sec}s  '
            f'dist={distance_m:.0f}m'
        )

    # ------------------------------------------------------------------
    def _build_cycles(self, session, points):
        cycles = []
        current = None
        last_cycle_id = 0

        for p in points:
            if p.state == STATE_IN_VEHICLE:
                if current:
                    cycles.append(current); current = None
            else:
                cid = p.cycle_id or (last_cycle_id + 1)
                if current is None or current['cycle_id'] != cid:
                    if current:
                        cycles.append(current)
                    last_cycle_id = cid
                    current = {'cycle_id': cid, 'start': p.recorded_at,
                               'end': p.recorded_at, 'points': [p]}
                else:
                    current['end'] = p.recorded_at
                    current['points'].append(p)
        if current:
            cycles.append(current)

        cyc_no = 0
        for c in cycles:
            pts = c['points']
            duration = (c['end'] - c['start']).total_seconds()
            if duration < 15 or len(pts) < 5:
                continue
            cyc_no += 1
            sr = sum(1 for x in pts if x.state == STATE_SEARCHING)
            dl = sum(1 for x in pts if x.state == STATE_DELIVERING)
            dist = 0.0
            for i in range(1, len(pts)):
                dist += haversine_m(pts[i - 1].lat, pts[i - 1].lon,
                                    pts[i].lat, pts[i].lon)
            avg_spd = sum(x.speed_kmh for x in pts) / len(pts)
            avg_acc = sum(x.accuracy_m for x in pts) / len(pts)

            cyc = Cycle.objects.create(
                session=session, cycle_no=cyc_no,
                started_at=c['start'], ended_at=c['end'],
                sr_seconds=sr, dl_seconds=dl, distance_m=dist,
                avg_speed_kmh=avg_spd, avg_accuracy_m=avg_acc,
            )
            LocationPoint.objects.filter(
                session=session,
                recorded_at__gte=c['start'], recorded_at__lte=c['end'],
            ).update(cycle_id=cyc.id)
