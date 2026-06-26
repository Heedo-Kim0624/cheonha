"""차량 현황_YYYY-MM-DD.xlsx → Vehicle 모델 일괄 import.

사용법:
    docker compose exec -T backend python manage.py import_vehicles_xlsx \
        /path/to/file.xlsx --company CHEONHA

엑셀 컬럼 (순서):
    차량번호 · VIN(TID) · 모델 · 차량 출고일 · 플릿 · 운전자 · 호기 · 배치현황 · 운영 구분
"""
import re
from datetime import date, datetime
from pathlib import Path
from openpyxl import load_workbook
from django.core.management.base import BaseCommand
from apps.vehicle_management.models import Company, Vehicle


PLACEMENT_MAP = {
    '운영중': 'OPERATING',
    '수리/대기중': 'REPAIRING',
    '수리': 'REPAIRING',
    '대기중': 'REPAIRING',
    '미출고': 'NOT_SHIPPED',
}
OPERATION_MAP = {
    '판매': 'SALE',
    '직영': 'DIRECT',
    '구독': 'SUBSCRIPTION',
    '기타': 'OTHER',
    '렌트': 'OTHER',
}


def _str(v):
    return ('' if v is None else str(v)).strip()


def _date(v):
    if isinstance(v, (date, datetime)):
        return v.date() if isinstance(v, datetime) else v
    s = _str(v)
    if not s:
        return None
    s = s.replace('.', '-').replace('/', '-')
    m = re.match(r'^(\d{4})-(\d{1,2})-(\d{1,2})', s)
    if not m:
        return None
    try:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None


def _short(num: str) -> str:
    digits = re.sub(r'\D', '', num)
    return digits[-4:] if len(digits) >= 4 else digits


class Command(BaseCommand):
    help = '차량현황 엑셀 import (헤더: 차량번호/VIN/모델/출고일/플릿/운전자/호기/배치현황/운영구분)'

    def add_arguments(self, parser):
        parser.add_argument('xlsx_path', type=str)
        parser.add_argument('--company', type=str, default='CHEONHA',
                            help='회사 코드 (YUHAN/CHEONHA/PERSONAL)')
        parser.add_argument('--sheet', type=str, default=None,
                            help='시트명 (기본: 첫 시트)')
        parser.add_argument('--dry-run', action='store_true')

    def handle(self, *args, **opts):
        path = Path(opts['xlsx_path']).expanduser()
        if not path.exists():
            self.stderr.write(f'파일 없음: {path}')
            return
        company = Company.objects.filter(code=opts['company']).first()
        if not company:
            self.stderr.write(f'회사 코드 {opts["company"]} 없음. Company.objects 확인.')
            return

        wb = load_workbook(path, data_only=True)
        ws = wb[opts['sheet']] if opts['sheet'] else wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            self.stderr.write('비어있는 시트')
            return

        header = [_str(v) for v in rows[0]]
        self.stdout.write(f'헤더: {header}')

        idx = {h: i for i, h in enumerate(header)}
        # 컬럼 인덱스 매핑 (한국어 헤더 매칭)
        def col(*candidates):
            for c in candidates:
                for h, i in idx.items():
                    if c in h:
                        return i
            return None
        c_num     = col('차량번호')
        c_vin     = col('VIN', 'TID')
        c_model   = col('모델')
        c_ship    = col('출고일')
        c_fleet   = col('플릿')
        c_driver  = col('운전자')
        c_hgi     = col('호기')
        c_place   = col('배치현황', '배치')
        c_op      = col('운영 구분', '운영구분', '운영')
        if c_num is None or c_vin is None:
            self.stderr.write(f'필수 컬럼 누락: 차량번호={c_num} VIN={c_vin}')
            return

        added, updated, skipped = 0, 0, 0
        for row in rows[1:]:
            num = _str(row[c_num])
            if not num:
                skipped += 1
                continue
            vin = _str(row[c_vin]) if c_vin is not None else ''
            data = {
                'company': company,
                'vehicle_number': num,
                'vehicle_number_short': _short(num),
                'vin_tid': vin,
                'model':   _str(row[c_model])  if c_model  is not None else '',
                'shipped_at': _date(row[c_ship]) if c_ship is not None else None,
                'fleet':   _str(row[c_fleet])  if c_fleet  is not None else '',
                'driver':  _str(row[c_driver]) if c_driver is not None else '',
                'hgi':     _str(row[c_hgi])    if c_hgi    is not None else '',
                'placement_status': PLACEMENT_MAP.get(_str(row[c_place]), 'OPERATING') if c_place is not None else 'OPERATING',
                'operation_type':   OPERATION_MAP.get(_str(row[c_op]), 'SALE')          if c_op    is not None else 'SALE',
            }
            if opts['dry_run']:
                self.stdout.write(f'  [{num}] {data["fleet"]} {data["hgi"]} {data["placement_status"]}')
                added += 1
                continue
            obj, created = Vehicle.objects.update_or_create(
                company=company, vehicle_number=num,
                defaults=data,
            )
            if created: added += 1
            else: updated += 1

        self.stdout.write(self.style.SUCCESS(
            f'완료: 신규 {added} / 갱신 {updated} / 스킵 {skipped} (회사 {company.name})'
        ))
