import hashlib
import calendar
from copy import deepcopy
from collections import defaultdict, deque
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.core.cache import cache
from django.core.files.base import ContentFile
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.db.models import Count, Max, Sum
from django.utils.dateparse import parse_date
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .constants import (
    ALLOWED_COMPANY_CODE,
    CATEGORY_COMMON,
    CATEGORY_LABELS,
    CATEGORY_SSG,
    CATEGORY_TRADERS,
    CATEGORY_YES24,
    JONGNO_CITY_NAME,
    N3_SSG_TABLE,
    SERVICE_LABELS,
    SERVICE_N3,
    SERVICE_W12,
    SERVICE_W4,
    SHIPPER_CODE,
    STATEMENT_ORDER,
    W12_SSG_TABLE,
    W12_VOLUME_BRACKETS,
    W4_TABLE,
    YES24_RATES,
)
from .models import OneDriver, OneDriverStatementOverride, OneSettlementUpload, OneShipmentOrder


REQUIRED_COLUMNS = ("일자", "주문번호", "수수료명칭", "시군구", "SM")
OPTIONAL_COLUMNS = ("운송장번호", "건별수수료", "최종처리시간", "받는분", "받는분주소", "비고")
SOURCE_TOTAL_COLUMN_OFFSET = 4
BLANK_ORDER_LOOKBACK = 10
RAW_SHEET_NAMES = ("RAW", "로우")
WATER_SHEET_NAME = "생수RAW"
WATER_SHEET_NAMES = (WATER_SHEET_NAME, "생수")
MANUAL_SHEET_NAME = "수기"
MANUAL_SHEET_NAMES = (MANUAL_SHEET_NAME,)
RAW_SHEET_FALLBACK_INDEX = 0
WATER_SHEET_FALLBACK_INDEX = 1
WATER_MANUAL_LABEL = "생수"
WATER_MANUAL_SOURCE = "one_water_raw"
WATER_ROW_AMOUNT = Decimal("200")
OFFSET_MANUAL_LABEL = "사고귀책 상계건"
OFFSET_MANUAL_SOURCE = "one_offset_items"
COLLECTION_PAGE_SIZE = 200
ONE_MONTH_CACHE_TTL = 300
COLLECTION_HEADERS = [
    ("date_label", "일자", "text"),
    ("order_number", "주문번호", "text"),
    ("waybill", "운송장번호", "text"),
    ("service_label", "구분", "text"),
    ("fee_label", "수수료명칭", "text"),
    ("city", "시군구", "text"),
    ("base_amount", "기준단가", "money"),
    ("volume_surcharge", "할증유형1(물량)", "money"),
    ("property_surcharge", "할증유형2(물성)", "money"),
    ("extra_surcharge", "할증유형3(기타)", "money"),
    ("total_amount", "합계", "money"),
    ("completed_at", "배송완료시간", "text"),
    ("driver_name", "SM", "text"),
    ("recipient", "받는분", "text"),
    ("address", "받는분주소", "text"),
    ("memo", "비고", "text"),
]

COLUMN_ALIASES = {
    "일자": ("배송일자", "배송 일자", "배송일", "배송 일", "작업일자", "작업 일자"),
    "주문번호": ("주문 번호", "오더번호", "오더 번호", "주문ID", "주문 ID"),
    "운송장번호": ("운송장 번호", "송장번호", "송장 번호", "운송장", "송장"),
    "수수료명칭": ("수수료 명칭", "수수료명", "수수료 명", "수수료"),
    "시군구": ("시/군/구", "시군구명", "시군구 명", "구", "지역"),
    "SM": ("sm", "기사", "기사명", "배송원", "배송원명", "배송원 명", "SM명", "SM 명"),
    "건별수수료": ("건별 수수료", "기준단가", "기준 단가", "수수료 금액", "금액"),
    "최종처리시간": ("최종 처리 시간", "배송완료시간", "배송 완료 시간", "완료시간", "완료 시간"),
    "받는분": ("받는 분", "수취인", "수령인", "고객명"),
    "받는분주소": ("받는분 주소", "받는 분 주소", "주소", "배송주소", "배송 주소"),
    "비고": ("메모", "특이사항", "특이 사항"),
    "배송일자": ("일자", "배송 일자", "배송일", "배송 일", "작업일자", "작업 일자"),
    "배송협력사": ("배송 협력사", "협력사", "배송사", "구분"),
    "품목명": ("상품명", "상품 명", "제품명", "제품 명", "품목"),
}
COLLECTION_SERVICE_LABELS = {
    SERVICE_W12: "W1,2",
    SERVICE_N3: "N3",
    SERVICE_W4: "W4",
}
COLLECTION_CATEGORY_LABELS = {
    CATEGORY_SSG: "SSG",
    CATEGORY_COMMON: "공동배송",
    CATEGORY_TRADERS: "트레이더스",
    CATEGORY_YES24: "YES24",
}
WEEKDAY_LABELS = ("월", "화", "수", "목", "금", "토", "일")
def _cache_part(value):
    if value is None:
        return "none"
    if hasattr(value, "timestamp"):
        return str(int(value.timestamp()))
    return str(value)


def _one_month_cache_key(kind, company_app, month):
    uploads = OneSettlementUpload.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
    ).aggregate(
        count=Count("id"),
        updated=Max("updated_at"),
        order_count=Sum("order_count"),
    )
    overrides = OneDriverStatementOverride.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
    ).aggregate(
        count=Count("id"),
        updated=Max("updated_at"),
    )
    return ":".join([
        "one",
        kind,
        str(company_app),
        str(month),
        _cache_part(uploads.get("count") or 0),
        _cache_part(uploads.get("updated")),
        _cache_part(uploads.get("order_count") or 0),
        _cache_part(overrides.get("count") or 0),
        _cache_part(overrides.get("updated")),
    ])


def _get_month_cache(kind, company_app, month):
    cached = cache.get(_one_month_cache_key(kind, company_app, month))
    return deepcopy(cached) if cached is not None else None


def _set_month_cache(kind, company_app, month, value):
    cache.set(_one_month_cache_key(kind, company_app, month), deepcopy(value), ONE_MONTH_CACHE_TTL)
    return value


DRIVER_ALIAS_MAP = {
    "\uc815\uc2b9\ud638": "\uc815\uc2b9\uc6a9",
    "\ucd5c\uae30\uc900": "\ucd5c\uae38\uc900",
}

# Compatibility exceptions found while reconciling the supplied 2026-05 ONE
# dashboard workbook. Keep these isolated so final ONE rules can replace them.
REFERENCE_ORDER_NUMBER_OVERRIDES = {
    ("2026-05-09", "조정훈", "6979-0724-9755"): "6979-0724-9755",
    ("2026-05-20", "박형배", "6981-1097-5740"): "6980-9960-0600",
}
REFERENCE_CATEGORY_OVERRIDES = {
    ("2026-05-16", "이현국", "10201307949"): (SERVICE_N3, CATEGORY_SSG),
}


@dataclass
class ParsedOrder:
    delivery_date: date
    order_number: str
    driver_name: str
    fee_name: str
    city: str
    boxes: int
    extra_boxes: int
    payload: dict
    service_code: str = ""
    category_code: str = ""
    is_mapped: bool = False
    base_amount: Decimal = Decimal("0")
    jongno_extra_amount: Decimal = Decimal("0")
    amount: Decimal = Decimal("0")


@dataclass
class ParsedManualItem:
    delivery_date: date
    driver_name: str
    label: str
    amount: Decimal
    payload: dict


def ensure_one_company_access(company_app):
    company_app = str(company_app or "").strip().lower()
    if company_app != ALLOWED_COMPANY_CODE:
        raise PermissionDenied("ONE 정산은 새회사 전용 기능입니다.")
    return company_app


def classify_fee_name(fee_name):
    value = str(fee_name or "").strip()
    upper = value.upper()

    # Temporary mapping. Replace this function when ONE confirms final fee labels.
    # Some ONE labels contain both N3 and 1W/2W/3W. N3 is the stronger signal.
    if ("당일" in value or "일요" in value) and "N3" in upper:
        if "YES24" in upper or "도서" in value:
            return SERVICE_N3, CATEGORY_YES24
        if "SSG" in upper:
            return SERVICE_N3, CATEGORY_SSG

    if "1W" in upper or "2W" in upper:
        if "YES24" in upper or "도서" in value:
            return SERVICE_W12, CATEGORY_YES24
        if "트레이더" in value or "애플" in value:
            return SERVICE_W12, CATEGORY_TRADERS
        if "공동" in value:
            return SERVICE_W12, CATEGORY_COMMON
        if "SSG" in upper:
            return SERVICE_W12, CATEGORY_SSG

    if "당일" in value or "일요" in value:
        if "YES24" in upper or "도서" in value:
            return SERVICE_W12, CATEGORY_YES24
        if "트레이더" in value or "애플" in value:
            return SERVICE_W12, CATEGORY_TRADERS
        if "공동" in value:
            return SERVICE_W12, CATEGORY_COMMON
        if "SSG" in upper:
            return SERVICE_W12, CATEGORY_SSG

    if "새벽" in value:
        if "공동배송" in value:
            return SERVICE_W4, CATEGORY_COMMON
        if "도서" in value or "YES24" in upper:
            return SERVICE_W4, CATEGORY_YES24
        if "SSG" in upper:
            return SERVICE_W4, CATEGORY_SSG

    return "", ""


def _parse_delivery_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    parsed = parse_date(str(value or "").strip())
    if not parsed:
        raise ValueError(f"일자 값을 해석할 수 없습니다: {value}")
    return parsed


def _cell_text(value):
    return "" if value is None else str(value).strip()


def _decimal_amount(value):
    text = _cell_text(value).replace(",", "")
    if not text or text == "-":
        return Decimal("0")
    try:
        return Decimal(text)
    except Exception:
        return Decimal("0")


def _driver_name_text(value):
    return _cell_text(value)


def _strip_driver_suffix(text):
    text = _cell_text(text)
    if text.endswith("_A"):
        text = text[:-2].strip()
    if len(text) > 1 and text.endswith("A") and not text[-2].isascii():
        text = text[:-1].strip()
    return text


def _manual_driver_name_candidates(value):
    text = _cell_text(value)
    candidates = []
    for candidate in (
        text,
        _strip_driver_suffix(text),
        DRIVER_ALIAS_MAP.get(text, ""),
        DRIVER_ALIAS_MAP.get(_strip_driver_suffix(text), ""),
    ):
        if candidate and candidate not in candidates:
            candidates.append(candidate)
    return candidates


def _match_manual_item_driver_names(manual_items, raw_driver_names):
    raw_driver_names = set(raw_driver_names)
    if not raw_driver_names:
        return manual_items
    for item in manual_items:
        for candidate in _manual_driver_name_candidates(item.driver_name):
            if candidate in raw_driver_names:
                item.driver_name = candidate
                break
    return manual_items


def _find_sheet_name(workbook, candidates):
    names_by_normalized = {
        _cell_text(sheet_name).casefold(): sheet_name
        for sheet_name in workbook.sheetnames
    }
    for candidate in candidates:
        sheet_name = names_by_normalized.get(_cell_text(candidate).casefold())
        if sheet_name:
            return sheet_name
    return ""


def _visible_non_empty_sheet_names(workbook):
    sheet_names = []
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        if getattr(worksheet, "sheet_state", "visible") != "visible":
            continue
        max_row = min(getattr(worksheet, "max_row", 20) or 20, 20)
        for row in worksheet.iter_rows(min_row=1, max_row=max_row, values_only=True):
            if row and any(_cell_text(value) for value in row):
                sheet_names.append(sheet_name)
                break
    return sheet_names


def _fallback_sheet_name(workbook, fallback_index):
    if fallback_index is None:
        return ""
    sheet_names = _visible_non_empty_sheet_names(workbook)
    if len(sheet_names) <= fallback_index:
        return ""
    return sheet_names[fallback_index]


def _sheet_rows(workbook, candidates, fallback_index=None):
    sheet_name = _find_sheet_name(workbook, candidates)
    if not sheet_name:
        sheet_name = _fallback_sheet_name(workbook, fallback_index)
    if not sheet_name:
        return "", []
    return sheet_name, list(workbook[sheet_name].iter_rows(values_only=True))


def _raw_sheet_rows(workbook):
    sheet_name, rows = _sheet_rows(workbook, RAW_SHEET_NAMES, RAW_SHEET_FALLBACK_INDEX)
    if not sheet_name:
        raise ValueError("RAW/로우 시트를 찾을 수 없습니다. 시트명이 다르면 첫 번째 시트를 RAW로 배치해 주세요.")
    return rows


def _raw_hash(rows):
    digest = hashlib.sha256()
    for row in rows:
        digest.update(repr(tuple(row)).encode("utf-8", "ignore"))
    return digest.hexdigest()


def _workbook_relevant_hash(workbook):
    digest = hashlib.sha256()
    for logical_name, candidates, fallback_index in (
        ("RAW", RAW_SHEET_NAMES, RAW_SHEET_FALLBACK_INDEX),
        (WATER_SHEET_NAME, WATER_SHEET_NAMES, WATER_SHEET_FALLBACK_INDEX),
        (MANUAL_SHEET_NAME, MANUAL_SHEET_NAMES, None),
    ):
        sheet_name = _find_sheet_name(workbook, candidates)
        if not sheet_name:
            sheet_name = _fallback_sheet_name(workbook, fallback_index)
        if not sheet_name:
            continue
        digest.update(logical_name.encode("utf-8", "ignore"))
        for row in workbook[sheet_name].iter_rows(values_only=True):
            digest.update(repr(tuple(row)).encode("utf-8", "ignore"))
    return digest.hexdigest()


def _normalized_column_name(value):
    text = _cell_text(value).casefold()
    return "".join(character for character in text if character.isalnum())


def _column_candidates(column_name):
    return (column_name, *COLUMN_ALIASES.get(column_name, ()))


def _build_header_index(row):
    exact_index = {}
    normalized_index = {}
    for position, value in enumerate(row):
        text = _cell_text(value)
        if not text:
            continue
        exact_index.setdefault(text, position)
        normalized = _normalized_column_name(text)
        if normalized:
            normalized_index.setdefault(normalized, position)

    header_index = dict(exact_index)
    for column_name in COLUMN_ALIASES:
        for candidate in _column_candidates(column_name):
            if candidate in exact_index:
                header_index[column_name] = exact_index[candidate]
                break
            normalized_candidate = _normalized_column_name(candidate)
            if normalized_candidate in normalized_index:
                header_index[column_name] = normalized_index[normalized_candidate]
                break
    return header_index


def _find_header_index(rows, required_columns, max_scan_rows=20):
    for row_index, row in enumerate(rows[:max_scan_rows]):
        index = _build_header_index(row)
        if all(name in index for name in required_columns):
            return index, row_index
    return {}, None


def _header_index(rows):
    if not rows:
        raise ValueError("RAW 시트가 비어 있습니다.")
    index, row_index = _find_header_index(rows, REQUIRED_COLUMNS)
    if not index:
        missing = ", ".join(REQUIRED_COLUMNS)
        raise ValueError(f"필수 컬럼이 없습니다: {missing}")
    return index, row_index


def _row_value(row, index, column_name):
    col = index.get(column_name)
    if col is None or col >= len(row):
        return None
    return row[col]


def _resolve_order_number(row, index, delivery_date, fee_name, driver_name, recent_order_keys):
    order_number = _cell_text(_row_value(row, index, "주문번호"))
    waybill = _cell_text(_row_value(row, index, "운송장번호"))
    completed_at = _cell_text(_row_value(row, index, "최종처리시간"))
    recipient = _cell_text(_row_value(row, index, "받는분"))
    address = _cell_text(_row_value(row, index, "받는분주소"))
    signature = (delivery_date, fee_name, driver_name, completed_at, recipient, address)

    if order_number:
        resolved = order_number
    else:
        matched = next(
            (
                item["order_number"]
                for item in reversed(recent_order_keys)
                if item["signature"] == signature
            ),
            "",
        )
        resolved = matched or waybill

    override = REFERENCE_ORDER_NUMBER_OVERRIDES.get((delivery_date.isoformat(), driver_name, waybill))
    if override:
        resolved = override

    if resolved:
        recent_order_keys.append({"signature": signature, "order_number": resolved})
    return resolved, order_number, waybill


def _w12_volume_rate(volume):
    volume = int(volume or 0)
    for threshold, rate in W12_VOLUME_BRACKETS:
        if threshold is None or volume < threshold:
            return rate
    return W12_VOLUME_BRACKETS[-1][1]


def _table_amount(boxes, table, extra_rate):
    boxes = max(1, int(boxes or 0))
    if boxes in table:
        return table[boxes]
    return table[max(table.keys())] + (Decimal(str(boxes - max(table.keys()))) * extra_rate)


def _w12_daily_volume_categories():
    return [CATEGORY_COMMON, CATEGORY_TRADERS]


def _calculate_one_base_amount(service_code, category_code, boxes, extra_boxes, delivery_date, w12_daily_rates):
    boxes = max(1, int(boxes or 0))
    extra_boxes = max(0, int(extra_boxes or 0))

    if service_code == SERVICE_W12 and category_code == CATEGORY_SSG:
        return _table_amount(boxes, W12_SSG_TABLE, Decimal("400"))
    if service_code == SERVICE_N3 and category_code == CATEGORY_SSG:
        return _table_amount(boxes, N3_SSG_TABLE, Decimal("400"))
    if service_code == SERVICE_W4 and category_code in {CATEGORY_SSG, CATEGORY_COMMON}:
        return _table_amount(boxes, W4_TABLE, Decimal("200"))
    if category_code == CATEGORY_YES24:
        return YES24_RATES.get(service_code, Decimal("0")) * boxes
    if service_code == SERVICE_W12 and category_code == CATEGORY_COMMON:
        return w12_daily_rates.get(delivery_date, Decimal("0")) * boxes
    if service_code == SERVICE_W12 and category_code == CATEGORY_TRADERS:
        return w12_daily_rates.get(delivery_date, Decimal("0")) + (Decimal(str(extra_boxes)) * Decimal("1000"))
    return Decimal("0")


def _apply_one_pricing_to_parsed_orders(parsed_orders):
    daily_w12_volume = defaultdict(int)
    for order in parsed_orders:
        if order.service_code == SERVICE_W12 and order.category_code in _w12_daily_volume_categories():
            daily_w12_volume[order.delivery_date] += 1

    w12_daily_rates = {
        delivery_date: _w12_volume_rate(volume)
        for delivery_date, volume in daily_w12_volume.items()
    }
    for order in parsed_orders:
        if not order.is_mapped:
            continue
        base = _calculate_one_base_amount(
            order.service_code,
            order.category_code,
            order.boxes,
            order.extra_boxes,
            order.delivery_date,
            w12_daily_rates,
        )
        order.base_amount = base
        order.jongno_extra_amount = Decimal("0")
        order.amount = base


def _classify_and_price_orders(parsed_orders):
    for order in parsed_orders:
        service_code, category_code = classify_fee_name(order.fee_name)
        override = REFERENCE_CATEGORY_OVERRIDES.get((
            order.delivery_date.isoformat(),
            order.driver_name,
            order.order_number,
        ))
        if override:
            service_code, category_code = override
        order.service_code = service_code
        order.category_code = category_code
        order.is_mapped = bool(service_code and category_code)
    _apply_one_pricing_to_parsed_orders(parsed_orders)


def recalculate_one_month_amounts(company_app, month):
    company_app = ensure_one_company_access(company_app)
    orders = OneShipmentOrder.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
        is_mapped=True,
    )
    daily_w12_volumes = dict(
        orders
        .filter(service_code=SERVICE_W12, category_code__in=_w12_daily_volume_categories())
        .values("delivery_date")
        .annotate(volume=Count("id"))
        .values_list("delivery_date", "volume")
    )
    w12_daily_rates = {
        delivery_date: _w12_volume_rate(volume)
        for delivery_date, volume in daily_w12_volumes.items()
    }

    changed = []
    for order in orders.only(
        "id",
        "delivery_date",
        "service_code",
        "category_code",
        "boxes",
        "extra_boxes",
        "amount",
        "base_amount",
        "jongno_extra_amount",
    ).iterator(chunk_size=2000):
        base = _calculate_one_base_amount(
            order.service_code,
            order.category_code,
            order.boxes,
            order.extra_boxes,
            order.delivery_date,
            w12_daily_rates,
        )
        if order.base_amount != base or order.jongno_extra_amount != Decimal("0") or order.amount != base:
            order.base_amount = base
            order.jongno_extra_amount = Decimal("0")
            order.amount = base
            changed.append(order)

    if changed:
        OneShipmentOrder.objects.bulk_update(
            changed,
            ["base_amount", "jongno_extra_amount", "amount"],
            batch_size=1000,
        )

    return {
        "month": month,
        "order_count": orders.count(),
        "updated_count": len(changed),
        "w12_daily_volume_days": len(daily_w12_volumes),
    }


def _sheet_header_index(rows, required_columns):
    if not rows:
        return {}, None
    return _find_header_index(rows, required_columns)


def _parse_water_manual_items(workbook):
    sheet_name, rows = _sheet_rows(workbook, WATER_SHEET_NAMES, WATER_SHEET_FALLBACK_INDEX)
    if not sheet_name:
        return []

    index, header_row_index = _sheet_header_index(rows, ("배송일자", "SM"))
    if not index:
        return []

    items = []
    for row_number, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
        if not row or not any(row):
            continue
        delivery_raw = _row_value(row, index, "배송일자")
        driver_name = _driver_name_text(_row_value(row, index, "SM"))
        if not delivery_raw or not driver_name:
            continue
        delivery_date = _parse_delivery_date(delivery_raw)

        payload = {
            "sheet": sheet_name,
            "row_number": row_number,
            "waybill": _cell_text(_row_value(row, index, "운송장번호")),
            "order_number": _cell_text(_row_value(row, index, "주문번호")),
            "item_name": _cell_text(_row_value(row, index, "품목명")),
            "address": _cell_text(_row_value(row, index, "받는분주소")),
        }
        items.append(ParsedManualItem(
            delivery_date=delivery_date,
            driver_name=driver_name,
            label=WATER_MANUAL_LABEL,
            amount=WATER_ROW_AMOUNT,
            payload=payload,
        ))
    return items


def parse_one_workbook_bundle(file_bytes):
    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    rows = _raw_sheet_rows(workbook)
    raw_hash = _workbook_relevant_hash(workbook)
    index, header_row_index = _header_index(rows)
    source_total_column = index.get("건별수수료", 6) + SOURCE_TOTAL_COLUMN_OFFSET
    grouped = {}
    total_rows = 0
    recent_order_keys = deque(maxlen=BLANK_ORDER_LOOKBACK)

    for row_number, row in enumerate(rows[header_row_index + 1:], start=header_row_index + 2):
        if not row or not any(row):
            continue
        delivery_raw = _row_value(row, index, "일자")
        driver_name = _driver_name_text(_row_value(row, index, "SM"))
        fee_name = _cell_text(_row_value(row, index, "수수료명칭"))
        if not delivery_raw or not driver_name or not fee_name:
            continue

        delivery_date = _parse_delivery_date(delivery_raw)
        order_number, original_order_number, waybill = _resolve_order_number(
            row,
            index,
            delivery_date,
            fee_name,
            driver_name,
            recent_order_keys,
        )
        if not order_number:
            continue
        city = _cell_text(_row_value(row, index, "시군구"))
        key = (delivery_date, fee_name, driver_name, order_number)
        total_rows += 1

        if key not in grouped:
            grouped[key] = {
                "delivery_date": delivery_date,
                "order_number": order_number,
                "driver_name": driver_name,
                "fee_name": fee_name,
                "city": city,
                "boxes": 0,
                "source_amount": Decimal("0"),
                "source_row_amounts": [],
                "row_numbers": [],
                "waybills": [],
                "sample": {},
            }
        item = grouped[key]
        item["boxes"] += 1
        row_source_amount = _decimal_amount(row[source_total_column] if len(row) > source_total_column else 0)
        item["source_amount"] += row_source_amount
        item["source_row_amounts"].append(row_source_amount)
        item["row_numbers"].append(row_number)
        if waybill:
            item["waybills"].append(waybill)
        if not original_order_number and waybill:
            item["sample"].setdefault("보정주문번호", order_number)
            item["sample"].setdefault("원본운송장번호", waybill)
        if not item["city"] and city:
            item["city"] = city
        if city == JONGNO_CITY_NAME:
            item["city"] = city
        for col in OPTIONAL_COLUMNS:
            value = _row_value(row, index, col)
            if value is not None and col not in item["sample"]:
                item["sample"][col] = _cell_text(value)

    parsed_orders = [
        ParsedOrder(
            delivery_date=item["delivery_date"],
            order_number=item["order_number"],
            driver_name=item["driver_name"],
            fee_name=item["fee_name"],
            city=item["city"],
            boxes=item["boxes"],
            extra_boxes=max(0, item["boxes"] - 1),
            payload={
                "row_numbers": item["row_numbers"],
                "waybills": item["waybills"],
                "source_amount": str(item["source_amount"]),
                "source_base_amount": str(item["source_row_amounts"][0] if item["source_row_amounts"] else Decimal("0")),
                "source_extra_amount": str(sum(item["source_row_amounts"][1:], Decimal("0"))),
                "source_row_amounts": [str(amount) for amount in item["source_row_amounts"]],
                "sample": item["sample"],
            },
        )
        for item in grouped.values()
    ]
    _classify_and_price_orders(parsed_orders)
    raw_driver_names = {order.driver_name for order in parsed_orders}
    manual_items = _match_manual_item_driver_names(_parse_water_manual_items(workbook), raw_driver_names)
    return raw_hash, total_rows, parsed_orders, manual_items


def parse_one_workbook(file_bytes):
    raw_hash, total_rows, parsed_orders, _manual_items = parse_one_workbook_bundle(file_bytes)
    return raw_hash, total_rows, parsed_orders


def _append_auto_manual_items(override, amounts_by_label, source):
    existing = override.manual_items or []
    retained = []
    current_amounts = defaultdict(int)
    for item in existing:
        if not isinstance(item, dict):
            continue
        if item.get("source") == source:
            label = str(item.get("label") or "").strip()
            current_amounts[label] += int(Decimal(str(item.get("amount") or 0)))
            continue
        retained.append(item)

    for label, amount in amounts_by_label.items():
        total = current_amounts[str(label)] + int(amount)
        if total:
            retained.append({
                "label": str(label),
                "amount": total,
                "source": source,
            })

    override.manual_items = retained
    override.save(update_fields=["manual_items", "updated_at"])


def rebuild_auto_manual_items(company_app, month):
    company_app = ensure_one_company_access(company_app)
    overrides = OneDriverStatementOverride.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
    )
    for override in overrides:
        retained = [
            item
            for item in (override.manual_items or [])
            if not (isinstance(item, dict) and item.get("source") == WATER_MANUAL_SOURCE)
        ]
        if retained != (override.manual_items or []):
            override.manual_items = retained
            override.save(update_fields=["manual_items", "updated_at"])

    manual_amounts = defaultdict(lambda: defaultdict(int))
    uploads = OneSettlementUpload.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
    ).order_by("id")
    for upload in uploads:
        if not upload.file:
            continue
        try:
            with upload.file.open("rb") as handle:
                _raw_hash, _total_rows, _parsed_orders, parsed_manual_items = parse_one_workbook_bundle(handle.read())
        except (FileNotFoundError, OSError, ValueError):
            continue
        for item in parsed_manual_items:
            manual_amounts[item.driver_name][item.label] += int(item.amount)

    if not manual_amounts:
        return

    drivers = {
        driver.name: driver
        for driver in OneDriver.objects.filter(company_app=company_app, name__in=manual_amounts.keys())
    }
    for driver_name, amounts in manual_amounts.items():
        driver = drivers.get(driver_name)
        if not driver:
            driver, _ = OneDriver.objects.get_or_create(company_app=company_app, name=driver_name)
            drivers[driver_name] = driver
        override = get_statement_override(company_app, month, driver)
        _append_auto_manual_items(override, amounts, WATER_MANUAL_SOURCE)


def import_one_file(*, company_app, uploaded_by, uploaded_file):
    company_app = ensure_one_company_access(company_app)
    original_filename = uploaded_file.name or ""
    file_bytes = uploaded_file.read()
    raw_hash, total_rows, parsed_orders, parsed_manual_items = parse_one_workbook_bundle(file_bytes)

    duplicate = OneSettlementUpload.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        raw_hash=raw_hash,
    ).first()
    if duplicate:
        return {"duplicate": True, "upload": duplicate}

    if not parsed_orders:
        raise ValueError("정산 대상 RAW 행이 없습니다.")

    dates = sorted({order.delivery_date for order in parsed_orders})
    month = dates[0].strftime("%Y-%m")
    unmapped_fee_names = sorted({order.fee_name for order in parsed_orders if not order.is_mapped})
    validation_errors = [
        {"type": "UNMAPPED_FEE_NAME", "fee_name": fee_name}
        for fee_name in unmapped_fee_names
    ]
    total_amount = sum((_order_source_amount(order) for order in parsed_orders), Decimal("0"))
    mapped_count = sum(1 for order in parsed_orders if order.is_mapped)

    with transaction.atomic():
        upload = OneSettlementUpload.objects.create(
            company_app=company_app,
            shipper_code=SHIPPER_CODE,
            month=month,
            delivery_date=dates[0],
            file=ContentFile(file_bytes, name=original_filename),
            original_filename=original_filename,
            raw_hash=raw_hash,
            total_rows=total_rows,
            order_count=len(parsed_orders),
            mapped_order_count=mapped_count,
            unmapped_order_count=len(parsed_orders) - mapped_count,
            total_amount=total_amount,
            validation_errors=validation_errors,
            status=OneSettlementUpload.Status.NEEDS_REVIEW if validation_errors else OneSettlementUpload.Status.IMPORTED,
            uploaded_by=uploaded_by if getattr(uploaded_by, "is_authenticated", False) else None,
        )

        drivers = {}
        driver_names = {order.driver_name for order in parsed_orders}
        driver_names.update(item.driver_name for item in parsed_manual_items)
        for name in sorted(driver_names):
            driver, _ = OneDriver.objects.get_or_create(company_app=company_app, name=name)
            drivers[name] = driver

        OneShipmentOrder.objects.bulk_create([
            OneShipmentOrder(
                upload=upload,
                company_app=company_app,
                shipper_code=SHIPPER_CODE,
                month=month,
                delivery_date=order.delivery_date,
                driver=drivers[order.driver_name],
                driver_name=order.driver_name,
                order_number=order.order_number,
                fee_name=order.fee_name,
                service_code=order.service_code,
                category_code=order.category_code,
                city=order.city,
                boxes=order.boxes,
                extra_boxes=order.extra_boxes,
                amount=order.amount,
                base_amount=order.base_amount,
                jongno_extra_amount=order.jongno_extra_amount,
                is_mapped=order.is_mapped,
                raw_payload=order.payload,
            )
            for order in parsed_orders
        ], batch_size=1000)

        manual_amounts = defaultdict(lambda: defaultdict(int))
        for item in parsed_manual_items:
            manual_amounts[item.driver_name][item.label] += int(item.amount)
        for driver_name, amounts in manual_amounts.items():
            override = get_statement_override(company_app, month, drivers[driver_name])
            _append_auto_manual_items(override, amounts, WATER_MANUAL_SOURCE)

    recalculate_one_month_amounts(company_app, month)
    return {"duplicate": False, "upload": upload}


def delete_one_upload(*, company_app, upload_id):
    company_app = ensure_one_company_access(company_app)
    upload = OneSettlementUpload.objects.get(
        id=upload_id,
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
    )
    month = upload.month
    file_name = upload.file.name if upload.file else ""
    storage = upload.file.storage if upload.file else None
    with transaction.atomic():
        upload.delete()
    if storage and file_name:
        storage.delete(file_name)
    recalculate_one_month_amounts(company_app, month)
    rebuild_auto_manual_items(company_app, month)
    return {"month": month}


def default_payment_due_date(month):
    year, month_number = [int(part) for part in str(month).split("-")]
    if month_number == 12:
        return date(year + 1, 1, 25)
    return date(year, month_number + 1, 25)


def month_bounds(month):
    year, month_number = [int(part) for part in str(month).split("-")]
    start = date(year, month_number, 1)
    end = date(year, month_number, calendar.monthrange(year, month_number)[1])
    return start, end


def _orders_for(company_app, month, driver_id=None):
    qs = OneShipmentOrder.objects.filter(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
        is_mapped=True,
    )
    if driver_id:
        qs = qs.filter(driver_id=driver_id)
    return qs


def _source_metric_from_payload(payload, field, fallback=0):
    if not isinstance(payload, dict):
        return int(Decimal(str(fallback or 0)))
    return int(_decimal_amount(payload.get(field, fallback)))


def _source_amount_from_payload(payload, fallback=0):
    return _source_metric_from_payload(payload, "source_amount", fallback)


def _source_row_amounts_from_payload(payload, fallback=0):
    if not isinstance(payload, dict):
        return [int(Decimal(str(fallback or 0)))]
    values = payload.get("source_row_amounts")
    if isinstance(values, list) and values:
        return [int(_decimal_amount(value)) for value in values]
    return [_source_amount_from_payload(payload, fallback)]


def _order_source_amount(order):
    payload = getattr(order, "payload", None)
    if payload is None:
        payload = getattr(order, "raw_payload", None)
    fallback = getattr(order, "amount", 0)
    return Decimal(str(_source_amount_from_payload(payload, fallback)))


def _manual_item_amount(item):
    if not isinstance(item, dict):
        return 0
    try:
        return int(Decimal(str(item.get("amount") or 0)))
    except (InvalidOperation, TypeError, ValueError):
        return 0


def _manual_item_is_offset(item):
    if not isinstance(item, dict):
        return False
    return item.get("source") == OFFSET_MANUAL_SOURCE or OFFSET_MANUAL_LABEL in str(item.get("label") or "")


def _clean_offset_rows(rows):
    cleaned = []
    if not isinstance(rows, list):
        return cleaned
    for row in rows:
        if not isinstance(row, dict):
            continue
        try:
            amount = int(Decimal(str(row.get("amount") or 0)))
        except (InvalidOperation, TypeError, ValueError):
            amount = 0
        cleaned.append({
            "waybill": str(row.get("waybill") or "").strip(),
            "reason": str(row.get("reason") or "").strip(),
            "amount": amount,
            "product": str(row.get("product") or "").strip(),
        })
    return cleaned


def _statement_offset_rows(statement):
    rows = []
    for item in statement.get("manual_items") or []:
        if not _manual_item_is_offset(item):
            continue
        rows.extend(_clean_offset_rows(item.get("rows") or []))
    return rows


def _statement_manual_totals_by_driver(company_app, month):
    totals = {}
    overrides = (
        OneDriverStatementOverride.objects
        .filter(company_app=company_app, shipper_code=SHIPPER_CODE, month=month)
        .select_related("driver")
        .only("driver_id", "driver__name", "manual_items")
    )
    for override in overrides:
        totals[override.driver_id] = {
            "driver_name": override.driver.name,
            "manual_total": sum(_manual_item_amount(item) for item in (override.manual_items or [])),
        }
    return totals


def _blank_metric():
    return {
        "households": 0,
        "boxes": 0,
        "extra_boxes": 0,
        "amount": 0,
        "source_amount": 0,
        "source_base_amount": 0,
        "source_extra_amount": 0,
    }


def _add_order_metric(bucket, order):
    bucket["households"] += 1
    bucket["boxes"] += int(order.boxes or 0)
    bucket["extra_boxes"] += int(order.extra_boxes or 0)
    bucket["amount"] += int(order.amount or 0)
    bucket["source_amount"] += _source_amount_from_payload(order.raw_payload, order.amount)
    bucket["source_base_amount"] += _source_metric_from_payload(order.raw_payload, "source_base_amount", order.amount)
    bucket["source_extra_amount"] += _source_metric_from_payload(order.raw_payload, "source_extra_amount", 0)


def _add_dashboard_group_metric(groups, key, order):
    if key not in groups:
        groups[key] = _blank_metric()
        groups[key]["households"] = 1
    bucket = groups[key]
    bucket["boxes"] += int(order.boxes or 0)
    bucket["amount"] += int(order.amount or 0)
    row_amounts = _source_row_amounts_from_payload(order.raw_payload, order.amount)
    bucket.setdefault("_source_row_amounts", []).extend(row_amounts)
    bucket["source_amount"] += sum(row_amounts)


def _finalize_dashboard_group_metric(row):
    row["extra_boxes"] = max(0, int(row["boxes"] or 0) - int(row["households"] or 0))
    row_amounts = row.get("_source_row_amounts") or []
    if row_amounts:
        row["source_base_amount"] = int(row_amounts[0] or 0)
        row["source_extra_amount"] = int(sum(row_amounts[1:]) or 0)
    return row


def _merge_metric(bucket, row):
    for field in ("households", "boxes", "extra_boxes", "amount", "source_amount", "source_base_amount", "source_extra_amount"):
        bucket[field] += int(row.get(field) or 0)


def build_month_summary(company_app, month):
    company_app = ensure_one_company_access(company_app)
    cached = _get_month_cache("summary", company_app, month)
    if cached is not None:
        return cached
    result = _build_month_summary_uncached(company_app, month)
    return _set_month_cache("summary", company_app, month, result)


def _build_month_summary_uncached(company_app, month):
    orders = _orders_for(company_app, month).order_by("delivery_date", "upload_id", "id").only(
        "id",
        "driver_id",
        "driver_name",
        "order_number",
        "delivery_date",
        "service_code",
        "category_code",
        "boxes",
        "extra_boxes",
        "amount",
        "raw_payload",
    )
    uploads = OneSettlementUpload.objects.filter(company_app=company_app, shipper_code=SHIPPER_CODE, month=month)
    aggregate = _blank_metric()
    by_service_map = defaultdict(_blank_metric)
    by_day_map = defaultdict(_blank_metric)
    by_day_service_map = defaultdict(_blank_metric)
    by_driver_service_map = defaultdict(_blank_metric)
    month_dashboard_groups = {}
    day_dashboard_groups = {}

    for order in orders.iterator(chunk_size=2000):
        order_number = str(order.order_number or order.id)
        _add_dashboard_group_metric(
            month_dashboard_groups,
            (order.service_code, order.category_code, order_number),
            order,
        )
        _add_dashboard_group_metric(
            day_dashboard_groups,
            (order.delivery_date, order.service_code, order.category_code, order_number),
            order,
        )
        _add_order_metric(
            by_driver_service_map[(order.driver_id, order.driver_name, order.service_code, order.category_code)],
            order,
        )

    for (service_code, category_code, _order_number), row in month_dashboard_groups.items():
        finalized = _finalize_dashboard_group_metric(row)
        _merge_metric(aggregate, finalized)
        _merge_metric(by_service_map[(service_code, category_code)], finalized)

    for (delivery_date, service_code, category_code, _order_number), row in day_dashboard_groups.items():
        finalized = _finalize_dashboard_group_metric(row)
        _merge_metric(by_day_map[delivery_date], finalized)
        _merge_metric(by_day_service_map[(delivery_date, service_code, category_code)], finalized)

    driver_ids = {driver_id for driver_id, _driver_name, _service_code, _category_code in by_driver_service_map}
    manual_totals = _statement_manual_totals_by_driver(company_app, month)
    manual_total = sum(int(manual_totals.get(driver_id, {}).get("manual_total") or 0) for driver_id in driver_ids)

    return {
        "month": month,
        "summary": {
            "upload_count": uploads.count(),
            "driver_count": len(driver_ids),
            "total_amount": int(aggregate["amount"] or 0),
            "manual_total": manual_total,
            "statement_total_amount": int(aggregate["amount"] or 0) + manual_total,
            "source_total_amount": int(aggregate["source_amount"] or 0),
            "source_base_amount": int(aggregate["source_base_amount"] or 0),
            "source_extra_amount": int(aggregate["source_extra_amount"] or 0),
            "households": int(aggregate["households"] or 0),
            "boxes": int(aggregate["boxes"] or 0),
            "extra_boxes": int(aggregate["extra_boxes"] or 0),
            "unmapped_order_count": int(sum(upload.unmapped_order_count for upload in uploads)),
        },
        "by_service": [
            {
                "service_code": service_code,
                "category_code": category_code,
                "service_label": SERVICE_LABELS.get(service_code, service_code),
                "category_label": CATEGORY_LABELS.get(category_code, category_code),
                "amount": int(row["amount"] or 0),
                "source_amount": int(row["source_amount"] or 0),
                "source_base_amount": int(row["source_base_amount"] or 0),
                "source_extra_amount": int(row["source_extra_amount"] or 0),
                "boxes": int(row["boxes"] or 0),
                "extra_boxes": int(row["extra_boxes"] or 0),
                "households": int(row["households"] or 0),
            }
            for (service_code, category_code), row in sorted(by_service_map.items())
        ],
        "by_day": [
            {
                "date": delivery_date.isoformat(),
                "households": int(row["households"] or 0),
                "boxes": int(row["boxes"] or 0),
                "extra_boxes": int(row["extra_boxes"] or 0),
                "amount": int(row["amount"] or 0),
                "source_amount": int(row["source_amount"] or 0),
                "source_base_amount": int(row["source_base_amount"] or 0),
                "source_extra_amount": int(row["source_extra_amount"] or 0),
            }
            for delivery_date, row in sorted(by_day_map.items())
        ],
        "by_day_service": [
            {
                "date": delivery_date.isoformat(),
                "service_code": service_code,
                "category_code": category_code,
                "service_label": SERVICE_LABELS.get(service_code, service_code),
                "category_label": CATEGORY_LABELS.get(category_code, category_code),
                "households": int(row["households"] or 0),
                "boxes": int(row["boxes"] or 0),
                "extra_boxes": int(row["extra_boxes"] or 0),
                "amount": int(row["amount"] or 0),
                "source_amount": int(row["source_amount"] or 0),
                "source_base_amount": int(row["source_base_amount"] or 0),
                "source_extra_amount": int(row["source_extra_amount"] or 0),
            }
            for (delivery_date, service_code, category_code), row in sorted(by_day_service_map.items())
        ],
        "by_driver_service": [
            {
                "driver_id": driver_id,
                "driver_name": driver_name,
                "service_code": service_code,
                "category_code": category_code,
                "service_label": SERVICE_LABELS.get(service_code, service_code),
                "category_label": CATEGORY_LABELS.get(category_code, category_code),
                "households": int(row["households"] or 0),
                "boxes": int(row["boxes"] or 0),
                "extra_boxes": int(row["extra_boxes"] or 0),
                "amount": int(row["amount"] or 0),
                "source_amount": int(row["source_amount"] or 0),
                "source_base_amount": int(row["source_base_amount"] or 0),
                "source_extra_amount": int(row["source_extra_amount"] or 0),
            }
            for (driver_id, driver_name, service_code, category_code), row in sorted(
                by_driver_service_map.items(),
                key=lambda item: (item[0][1], item[0][2], item[0][3]),
            )
        ],
    }


def list_drivers(company_app, month):
    company_app = ensure_one_company_access(company_app)
    cached = _get_month_cache("drivers", company_app, month)
    if cached is not None:
        return cached
    summary = build_month_summary(company_app, month)
    grouped = defaultdict(_blank_metric)
    for row in summary.get("by_driver_service") or []:
        key = (row["driver_id"], row["driver_name"])
        _merge_metric(grouped[key], row)
    manual_totals = _statement_manual_totals_by_driver(company_app, month)
    result = [
        {
            "id": driver_id,
            "name": driver_name,
            "households": int(row["households"] or 0),
            "boxes": int(row["boxes"] or 0),
            "extra_boxes": int(row["extra_boxes"] or 0),
            "amount": int(row["amount"] or 0),
            "manual_total": int(manual_totals.get(driver_id, {}).get("manual_total") or 0),
            "statement_total_amount": int(row["amount"] or 0) + int(manual_totals.get(driver_id, {}).get("manual_total") or 0),
            "source_amount": int(row["source_amount"] or 0),
        }
        for (driver_id, driver_name), row in sorted(grouped.items(), key=lambda item: item[0][1])
    ]
    return _set_month_cache("drivers", company_app, month, result)


def _collection_orders(company_app, month):
    return (
        OneShipmentOrder.objects
        .filter(company_app=company_app, shipper_code=SHIPPER_CODE, month=month)
        .only(
            "id",
            "delivery_date",
            "driver_name",
            "order_number",
            "raw_payload",
            "boxes",
            "base_amount",
            "jongno_extra_amount",
            "amount",
            "category_code",
            "service_code",
            "fee_name",
            "city",
            "is_mapped",
            "extra_boxes",
        )
        .order_by("delivery_date", "driver_name", "service_code", "category_code", "order_number", "id")
    )


def _format_collection_date(value):
    if not value:
        return ""
    return f"{value.strftime('%y-%m-%d')}-{WEEKDAY_LABELS[value.weekday()]}"


def _split_amount(total, count):
    count = max(1, int(count or 1))
    amount = int(Decimal(str(total or 0)))
    sign = -1 if amount < 0 else 1
    base, remainder = divmod(abs(amount), count)
    return [sign * (base + (1 if index < remainder else 0)) for index in range(count)]


def _collection_fee_label(order):
    if order.category_code == CATEGORY_COMMON and order.service_code == SERVICE_W4:
        return "공동배송망"
    if order.category_code:
        return COLLECTION_CATEGORY_LABELS.get(order.category_code, order.category_code)
    return order.fee_name


def _collection_memo(order, sample):
    memo = str(sample.get("비고") or "").strip()
    if memo:
        return memo
    if order.is_mapped:
        return f"착지:{order.boxes}/ 추가박스:{order.extra_boxes}"
    return "미매핑 수수료명칭"


def _expanded_collection_rows(order):
    payload = order.raw_payload or {}
    sample = payload.get("sample") or {}
    waybills = [str(item or "").strip() for item in (payload.get("waybills") or []) if str(item or "").strip()]
    if not waybills:
        fallback_waybill = str(sample.get("운송장번호") or sample.get("원본운송장번호") or "").strip()
        waybills = [fallback_waybill] if fallback_waybill else []

    row_count = max(1, int(order.boxes or 0), len(waybills))
    waybills = waybills + [""] * max(0, row_count - len(waybills))
    base_amounts = _split_amount(order.base_amount, row_count)
    extra_amounts = _split_amount(order.jongno_extra_amount, row_count)
    total_amounts = _split_amount(order.amount, row_count)
    service_label = COLLECTION_SERVICE_LABELS.get(order.service_code, order.service_code or "")
    fee_label = _collection_fee_label(order)
    memo = _collection_memo(order, sample)

    for index in range(row_count):
        yield {
            "date_label": _format_collection_date(order.delivery_date),
            "order_number": order.order_number,
            "waybill": waybills[index],
            "service_label": service_label,
            "fee_label": fee_label,
            "city": order.city,
            "base_amount": base_amounts[index],
            "volume_surcharge": 0,
            "property_surcharge": 0,
            "extra_surcharge": extra_amounts[index],
            "total_amount": total_amounts[index],
            "completed_at": str(sample.get("최종처리시간") or "").strip(),
            "driver_name": order.driver_name,
            "recipient": str(sample.get("받는분") or "").strip(),
            "address": str(sample.get("받는분주소") or "").strip(),
            "memo": memo,
        }


def _collection_missing_order_count(company_app, month):
    company_app = ensure_one_company_access(company_app)
    cached = _get_month_cache("collection_missing", company_app, month)
    if cached is not None:
        return int(cached)
    missing = 0
    orders = (
        OneShipmentOrder.objects
        .filter(company_app=company_app, shipper_code=SHIPPER_CODE, month=month)
        .only("raw_payload", "boxes")
        .iterator(chunk_size=2000)
    )
    for order in orders:
        sample = (order.raw_payload or {}).get("sample") or {}
        if sample.get("보정주문번호"):
            missing += max(1, int(order.boxes or 0))
    _set_month_cache("collection_missing", company_app, month, missing)
    return missing


def _collection_columns(missing_order_count):
    return [
        {
            "key": key,
            "label": f"주문번호 (누락: {missing_order_count}건)" if key == "order_number" else label,
            "type": column_type,
        }
        for key, label, column_type in COLLECTION_HEADERS
    ]


def build_collection_table(company_app, month, page=1, page_size=COLLECTION_PAGE_SIZE):
    company_app = ensure_one_company_access(company_app)
    page = max(1, int(page or 1))
    page_size = max(50, min(500, int(page_size or COLLECTION_PAGE_SIZE)))
    cache_kind = f"collection_page:{page}:{page_size}"
    cached = _get_month_cache(cache_kind, company_app, month)
    if cached is not None:
        return cached
    offset = (page - 1) * page_size
    orders = _collection_orders(company_app, month)
    total_rows = int(orders.aggregate(total=Sum("boxes"))["total"] or 0)
    missing_order_count = _collection_missing_order_count(company_app, month)

    rows = []
    seen = 0
    for order in orders:
        for row in _expanded_collection_rows(order):
            if seen >= offset and len(rows) < page_size:
                rows.append(row)
            seen += 1
            if len(rows) >= page_size and seen >= offset + page_size:
                break
        if len(rows) >= page_size and seen >= offset + page_size:
            break

    result = {
        "month": month,
        "columns": _collection_columns(missing_order_count),
        "rows": rows,
        "page": page,
        "page_size": page_size,
        "total": total_rows,
        "missing_order_count": missing_order_count,
    }
    return _set_month_cache(cache_kind, company_app, month, result)


def build_collection_excel(company_app, month):
    company_app = ensure_one_company_access(company_app)
    orders = _collection_orders(company_app, month)
    missing_order_count = _collection_missing_order_count(company_app, month)
    columns = _collection_columns(missing_order_count)

    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("취합B_(기사)")
    worksheet.append([column["label"] for column in columns])
    for order in orders:
        for row in _expanded_collection_rows(order):
            worksheet.append([row.get(column["key"], "") for column in columns])

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()


def get_statement_override(company_app, month, driver):
    override, _ = OneDriverStatementOverride.objects.get_or_create(
        company_app=company_app,
        shipper_code=SHIPPER_CODE,
        month=month,
        driver=driver,
        defaults={"payment_due_date": default_payment_due_date(month)},
    )
    if not override.payment_due_date:
        override.payment_due_date = default_payment_due_date(month)
        override.save(update_fields=["payment_due_date", "updated_at"])
    return override


def build_driver_statement(company_app, month, driver_id):
    company_app = ensure_one_company_access(company_app)
    driver = OneDriver.objects.get(company_app=company_app, id=driver_id)
    override = get_statement_override(company_app, month, driver)
    orders = _orders_for(company_app, month, driver_id=driver.id)

    grouped = {
        key: {"service_code": key[0], "category_code": key[1], "households": 0, "extra_boxes": 0, "boxes": 0, "amount": 0}
        for key in STATEMENT_ORDER
    }
    raw_grouped = orders.values("service_code", "category_code").annotate(
        households=Count("id"),
        boxes=Sum("boxes"),
        extra_boxes=Sum("extra_boxes"),
        amount=Sum("amount"),
    )
    for row in raw_grouped:
        key = (row["service_code"], row["category_code"])
        if key not in grouped:
            grouped[key] = {"service_code": key[0], "category_code": key[1], "households": 0, "extra_boxes": 0, "boxes": 0, "amount": 0}
        grouped[key].update({
            "households": int(row["households"] or 0),
            "boxes": int(row["boxes"] or 0),
            "extra_boxes": int(row["extra_boxes"] or 0),
            "amount": int(row["amount"] or 0),
        })

    daily = defaultdict(lambda: defaultdict(lambda: {"households": 0, "extra_boxes": 0, "boxes": 0, "amount": 0}))
    for row in orders.values("delivery_date", "service_code", "category_code").annotate(
        households=Count("id"),
        boxes=Sum("boxes"),
        extra_boxes=Sum("extra_boxes"),
        amount=Sum("amount"),
    ):
        bucket = daily[row["delivery_date"]][f"{row['service_code']}:{row['category_code']}"]
        bucket["households"] = int(row["households"] or 0)
        bucket["boxes"] = int(row["boxes"] or 0)
        bucket["extra_boxes"] = int(row["extra_boxes"] or 0)
        bucket["amount"] = int(row["amount"] or 0)

    summary_rows = []
    for service_code, category_code in STATEMENT_ORDER:
        row = grouped[(service_code, category_code)]
        summary_rows.append({
            **row,
            "service_label": SERVICE_LABELS.get(service_code, service_code),
            "category_label": CATEGORY_LABELS.get(category_code, category_code),
        })

    manual_items = override.manual_items or []
    manual_total = sum(_manual_item_amount(item) for item in manual_items)
    calculated_total = sum(row["amount"] for row in summary_rows)
    total_amount = calculated_total + manual_total
    start, end = month_bounds(month)
    return {
        "id": override.id,
        "driver": {"id": driver.id, "name": driver.name},
        "month": month,
        "period_start": start.isoformat(),
        "period_end": end.isoformat(),
        "payment_due_date": override.payment_due_date.isoformat(),
        "memo": override.memo,
        "manual_items": manual_items,
        "summary_rows": summary_rows,
        "daily_rows": [
            {
                "date": d.isoformat(),
                "values": dict(daily[d]),
            }
            for d in sorted(daily)
        ],
        "calculated_total": calculated_total,
        "manual_total": manual_total,
        "total_amount": total_amount,
    }


def update_statement_override(*, override_id, company_app, payload, user):
    company_app = ensure_one_company_access(company_app)
    override = OneDriverStatementOverride.objects.select_related("driver").get(id=override_id, company_app=company_app)
    if "payment_due_date" in payload:
        override.payment_due_date = parse_date(str(payload.get("payment_due_date") or "")) or default_payment_due_date(override.month)
    if "memo" in payload:
        override.memo = str(payload.get("memo") or "")
    if "manual_items" in payload:
        items = payload.get("manual_items") or []
        cleaned = []
        for item in items:
            if not isinstance(item, dict):
                continue
            label = str(item.get("label") or "").strip()
            if not label:
                continue
            amount = int(Decimal(str(item.get("amount") or 0)))
            cleaned_item = {"label": label, "amount": amount}
            source = str(item.get("source") or "").strip()
            if source:
                cleaned_item["source"] = source
            rows = _clean_offset_rows(item.get("rows") or [])
            if rows:
                cleaned_item["rows"] = rows
                if source == OFFSET_MANUAL_SOURCE or OFFSET_MANUAL_LABEL in label:
                    cleaned_item["source"] = OFFSET_MANUAL_SOURCE
                    cleaned_item["label"] = OFFSET_MANUAL_LABEL
                    cleaned_item["amount"] = sum(row["amount"] for row in rows)
            cleaned.append(cleaned_item)
        override.manual_items = cleaned
    if getattr(user, "is_authenticated", False):
        override.updated_by = user
    override.save()
    return override


def _statement_title(statement):
    year, month_number = [int(part) for part in statement["month"].split("-")]
    return f"{year}년 {month_number}월 운송료 지급명세서"


def build_statement_excel(statement):
    wb = Workbook()
    ws = wb.active
    ws.title = "지급명세서"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.print_area = "A1:O60"

    black_fill = PatternFill("solid", fgColor="000000")
    gray_fill = PatternFill("solid", fgColor="D9D9D9")
    white_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    title_font = Font(bold=True, size=18)
    month_font = Font(bold=True, size=16)
    thin = Side(style="thin", color="000000")
    medium = Side(style="medium", color="000000")
    thin_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    right = Alignment(horizontal="right", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    def number(value):
        return int(value or 0)

    def compact_date(value):
        parsed = parse_date(str(value or ""))
        if not parsed:
            return "-"
        return f"{str(parsed.year)[2:]}-{parsed.month:02d}-{parsed.day:02d}"

    def short_date(value):
        parsed = parse_date(str(value or ""))
        if not parsed:
            return str(value or "-")
        weekday_labels = ("월", "화", "수", "목", "금", "토", "일")
        return f"{compact_date(value)}-{weekday_labels[parsed.weekday()]}"

    def korean_date(value):
        parsed = parse_date(str(value or ""))
        if not parsed:
            return "-"
        return f"{parsed.year}년 {parsed.month}월 {parsed.day}일"

    def extra_box_value(row):
        if row.get("category_code") == CATEGORY_YES24:
            return "-"
        if row.get("service_code") == SERVICE_W12 and row.get("category_code") == CATEGORY_COMMON:
            return "-"
        if row.get("service_code") == SERVICE_N3 and row.get("category_code") == CATEGORY_YES24:
            return "-"
        return number(row.get("extra_boxes"))

    def manual_amount(keyword):
        total = 0
        for item in statement.get("manual_items") or []:
            if keyword in str(item.get("label") or ""):
                total += int(item.get("amount") or 0)
        return total

    def daily_value(row, key, field):
        return number((row.get("values") or {}).get(key, {}).get(field) or 0)

    def daily_rows_for_month():
        start = parse_date(statement.get("period_start") or "")
        end = parse_date(statement.get("period_end") or "")
        if not start or not end:
            return []
        values_by_date = {
            item.get("date"): item.get("values") or {}
            for item in statement.get("daily_rows") or []
        }
        rows = []
        cursor = start
        while cursor <= end:
            iso = cursor.isoformat()
            rows.append({"date": iso, "values": values_by_date.get(iso, {})})
            cursor = date.fromordinal(cursor.toordinal() + 1)
        return rows

    month_parts = str(statement.get("month") or "").split("-")
    year = month_parts[0] if month_parts else ""
    month_number = int(month_parts[1]) if len(month_parts) > 1 and month_parts[1].isdigit() else 0
    month_label = f"{year}년 {month_number}월"
    period_text = f"{compact_date(statement.get('period_start'))} 부터 {compact_date(statement.get('period_end'))} 까지"

    for col in range(1, 16):
        ws.column_dimensions[get_column_letter(col)].width = 10.75 if col == 1 else 13
    for row in range(1, 61):
        ws.row_dimensions[row].height = 18
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30
    ws.row_dimensions[8].height = 10
    ws.row_dimensions[26].height = 10

    merges = [
        "A1:B1", "A2:B2", "B4:C4", "B5:C5", "B6:C6", "B7:C7",
        "A9:I9", "B10:C10", "D10:E10", "F10:G10", "H10:I10",
        "B11:B14", "B15:B16", "B17:B19",
        "D11:E11", "F11:G11", "H11:I11",
        "D12:E12", "F12:G12", "H12:I12",
        "D13:E13", "F13:G13", "H13:I13",
        "D14:E14", "F14:G14", "H14:I14",
        "D15:E15", "F15:G15", "H15:I15",
        "D16:E16", "F16:G16", "H16:I16",
        "D17:E17", "F17:G17", "H17:I17",
        "D18:E18", "F18:G18", "H18:I18",
        "D19:E19", "F19:G19", "H19:I19",
        "D20:E20", "F20:G20", "H20:I20",
        "B21:B22", "D21:E21", "F21:G21", "H21:I21",
        "D22:E22", "F22:G22", "H22:I22",
        "A23:C23", "H23:I23",
        "A24:I24", "A25:C25", "H25:I25",
        "A27:A29", "C27:G27", "I27:J27", "L27:O27",
        "B28:C28", "D28:E28", "H28:I28", "K28:L28", "M28:N28",
    ]

    def set_cell(ref, value, font=None, fill=None, alignment=None, number_format=None):
        cell = ws[ref]
        cell.value = value
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if alignment:
            cell.alignment = alignment
        if number_format:
            cell.number_format = number_format

    set_cell("A1", month_label, month_font, alignment=left)
    set_cell("A2", "운송료 지급명세서", title_font, alignment=left)
    set_cell("A4", "수수료 명.", bold, alignment=left)
    set_cell("B4", "CJ 대한통운 ONE 배송건", bold, alignment=left)
    set_cell("A5", "지급 대상자.", bold, alignment=left)
    set_cell("B5", statement["driver"]["name"], bold, alignment=left)
    set_cell("A6", "정산 기간일.", bold, alignment=left)
    set_cell("B6", period_text, bold, alignment=left)
    set_cell("A7", "지급 예정일.", bold, alignment=left)
    set_cell("B7", korean_date(statement.get("payment_due_date")), bold, alignment=left)

    offset_rows = _statement_offset_rows(statement)
    for index, row in enumerate(offset_rows[:24], start=1):
        set_cell(f"J{index}", row["waybill"], alignment=left)
        set_cell(f"K{index}", row["reason"], alignment=left)
        set_cell(f"L{index}", row["amount"], alignment=right, number_format="#,##0")
        set_cell(f"M{index}", row["product"], alignment=left)
    if len(offset_rows) > 24:
        set_cell("J25", f"외 {len(offset_rows) - 24:,}건", bold, alignment=left)
    if offset_rows:
        set_cell("L26", sum(row["amount"] for row in offset_rows), bold, alignment=right, number_format="#,##0")

    set_cell("A9", "지급내역", white_font, black_fill, center)
    set_cell("B10", "구 분", bold, alignment=center)
    set_cell("D10", "착지 / 건", bold, alignment=center)
    set_cell("F10", "추가박스", bold, alignment=center)
    set_cell("H10", "금액 (부가세 포함)", bold, alignment=center)
    set_cell("A11", "전체 운송료", bold, alignment=center)

    summary_lookup = {
        (item["service_code"], item["category_code"]): item
        for item in statement["summary_rows"]
    }
    template_order = [
        (SERVICE_W12, CATEGORY_SSG),
        (SERVICE_W12, CATEGORY_COMMON),
        (SERVICE_W12, CATEGORY_TRADERS),
        (SERVICE_W12, CATEGORY_YES24),
        (SERVICE_N3, CATEGORY_SSG),
        (SERVICE_N3, CATEGORY_YES24),
        (SERVICE_W4, CATEGORY_SSG),
        (SERVICE_W4, CATEGORY_COMMON),
        (SERVICE_W4, CATEGORY_YES24),
    ]
    service_first_labels = {0: "W1,2_당일", 4: "N3_당일", 6: "W4_새벽"}
    for index, (service_code, category_code) in enumerate(template_order):
        row = 11 + index
        item = summary_lookup.get((service_code, category_code), {
            "service_code": service_code,
            "category_code": category_code,
            "category_label": CATEGORY_LABELS.get(category_code, category_code),
            "households": 0,
            "extra_boxes": 0,
            "amount": 0,
        })
        set_cell(f"B{row}", service_first_labels.get(index, ""), bold if index in service_first_labels else None, alignment=center)
        set_cell(f"C{row}", item["category_label"], alignment=left)
        set_cell(f"D{row}", number(item["households"]), alignment=right, number_format="#,##0")
        set_cell(f"F{row}", extra_box_value(item), alignment=right, number_format="#,##0")
        set_cell(f"H{row}", number(item["amount"]), alignment=right, number_format="#,##0")

    set_cell("B20", "통 합", bold, alignment=center)
    set_cell("C20", "생수", alignment=left)
    set_cell("D20", 0, alignment=right, number_format="#,##0")
    set_cell("F20", 0, alignment=right, number_format="#,##0")
    set_cell("H20", manual_amount("생수"), alignment=right, number_format="#,##0")
    set_cell("B21", "운송료 외", bold, alignment=center)
    set_cell("C21", "추가 전달", alignment=left)
    set_cell("D21", 0, alignment=right, number_format="#,##0")
    set_cell("F21", "-", alignment=right)
    set_cell("H21", manual_amount("추가 전달"), alignment=right, number_format="#,##0")
    set_cell("C22", "수당", alignment=left)
    set_cell("D22", "-", alignment=right)
    set_cell("F22", "-", alignment=right)
    set_cell("H22", manual_amount("수당"), alignment=right, number_format="#,##0")
    set_cell("A23", "사고귀책 상계건", bold, gray_fill, center)
    set_cell("G23", "-", alignment=right)
    set_cell("H23", manual_amount("사고귀책"), alignment=right, number_format="#,##0")
    set_cell("A24", "", fill=black_fill)
    set_cell("A25", "지급 합계", bold, alignment=center)
    set_cell("H25", number(statement["total_amount"]), bold, alignment=right, number_format="#,##0")

    set_cell("A27", "일 자", white_font, black_fill, center)
    set_cell("B27", "구 분", bold, alignment=center)
    set_cell("C27", "W1.2_당일", white_font, black_fill, center)
    set_cell("H27", "구 분", bold, alignment=center)
    set_cell("I27", "N3_당일", white_font, black_fill, center)
    set_cell("K27", "구 분", bold, alignment=center)
    set_cell("L27", "4W_새벽", white_font, black_fill, center)
    for ref, value in {
        "B28": "SSG", "D28": "트레이더스", "F28": "공동배송", "G28": "YES24",
        "H28": "SSG", "J28": "YES24", "K28": "SSG", "M28": "공동배송", "O28": "YES24",
    }.items():
        set_cell(ref, value, bold, alignment=center)
    for col, value in enumerate([
        "착지/건", "추가박스", "착지/건", "추가박스", "착지/건", "착지 / 건",
        "착지/건", "추가박스", "착지 / 건", "착지/건", "추가박스", "착지/건",
        "추가박스", "착지 / 건",
    ], start=2):
        set_cell(f"{get_column_letter(col)}29", value, bold, alignment=center)

    for daily in daily_rows_for_month():
        parsed = parse_date(daily["date"])
        if not parsed:
            continue
        row = 29 + parsed.day
        if not 30 <= row <= 60:
            continue
        set_cell(f"A{row}", short_date(daily["date"]), alignment=center)
        values = [
            daily_value(daily, f"{SERVICE_W12}:{CATEGORY_SSG}", "households"),
            daily_value(daily, f"{SERVICE_W12}:{CATEGORY_SSG}", "extra_boxes"),
            daily_value(daily, f"{SERVICE_W12}:{CATEGORY_TRADERS}", "households"),
            daily_value(daily, f"{SERVICE_W12}:{CATEGORY_TRADERS}", "extra_boxes"),
            daily_value(daily, f"{SERVICE_W12}:{CATEGORY_COMMON}", "households"),
            daily_value(daily, f"{SERVICE_W12}:{CATEGORY_YES24}", "households"),
            daily_value(daily, f"{SERVICE_N3}:{CATEGORY_SSG}", "households"),
            daily_value(daily, f"{SERVICE_N3}:{CATEGORY_SSG}", "extra_boxes"),
            daily_value(daily, f"{SERVICE_N3}:{CATEGORY_YES24}", "households"),
            daily_value(daily, f"{SERVICE_W4}:{CATEGORY_SSG}", "households"),
            daily_value(daily, f"{SERVICE_W4}:{CATEGORY_SSG}", "extra_boxes"),
            daily_value(daily, f"{SERVICE_W4}:{CATEGORY_COMMON}", "households"),
            daily_value(daily, f"{SERVICE_W4}:{CATEGORY_COMMON}", "extra_boxes"),
            daily_value(daily, f"{SERVICE_W4}:{CATEGORY_YES24}", "households"),
        ]
        for col, value in enumerate(values, start=2):
            set_cell(f"{get_column_letter(col)}{row}", value, alignment=right, number_format="#,##0")

    for row in range(9, 26):
        for col in range(1, 10):
            ws.cell(row=row, column=col).border = thin_border
    for row in range(27, 61):
        for col in range(1, 16):
            ws.cell(row=row, column=col).border = thin_border
    for row in (15, 17):
        for col in range(2, 10):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(
                left=cell.border.left,
                right=cell.border.right,
                top=medium,
                bottom=cell.border.bottom,
            )
    for col in range(1, 10):
        ws.cell(row=23, column=col).fill = gray_fill
        ws.cell(row=24, column=col).fill = black_fill
    for ref in ("A27", "C27", "I27", "L27"):
        ws[ref].fill = black_fill
        ws[ref].font = white_font

    for merge_range in merges:
        ws.merge_cells(merge_range)

    for row in ws.iter_rows(min_row=1, max_row=60, min_col=1, max_col=15):
        for cell in row:
            if cell.alignment == Alignment():
                cell.alignment = Alignment(vertical="center")

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def build_statement_pdf(statement):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    font_name = "Helvetica"
    bold_font_name = "Helvetica-Bold"
    for font_path in (
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc",
        "C:/Windows/Fonts/malgun.ttf",
    ):
        try:
            pdfmetrics.registerFont(TTFont("OneKorean", font_path))
            font_name = "OneKorean"
            break
        except Exception:
            continue
    for font_path in (
        "/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf",
        "/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc",
        "C:/Windows/Fonts/malgunbd.ttf",
    ):
        try:
            pdfmetrics.registerFont(TTFont("OneKoreanBold", font_path))
            bold_font_name = "OneKoreanBold"
            break
        except Exception:
            bold_font_name = font_name

    def number(value):
        return f"{int(value or 0):,}"

    def compact_date(value):
        parsed = parse_date(str(value or ""))
        if not parsed:
            return "-"
        return f"{str(parsed.year)[2:]}-{parsed.month:02d}-{parsed.day:02d}"

    def short_date(value):
        parsed = parse_date(str(value or ""))
        if not parsed:
            return str(value or "-")
        weekday_labels = ("월", "화", "수", "목", "금", "토", "일")
        return f"{compact_date(value)}-{weekday_labels[parsed.weekday()]}"

    def korean_date(value):
        parsed = parse_date(str(value or ""))
        if not parsed:
            return "-"
        return f"{parsed.year}년 {parsed.month}월 {parsed.day}일"

    def extra_box_value(row):
        if row.get("category_code") == CATEGORY_YES24:
            return "-"
        if row.get("service_code") == SERVICE_W12 and row.get("category_code") == CATEGORY_COMMON:
            return "-"
        if row.get("service_code") == SERVICE_N3 and row.get("category_code") == CATEGORY_YES24:
            return "-"
        return number(row.get("extra_boxes"))

    def manual_amount(keyword):
        total = 0
        for item in statement.get("manual_items") or []:
            if keyword in str(item.get("label") or ""):
                total += int(item.get("amount") or 0)
        return total

    def daily_value(row, key, field):
        return number((row.get("values") or {}).get(key, {}).get(field) or 0)

    def daily_rows_for_month():
        start = parse_date(statement.get("period_start") or "")
        end = parse_date(statement.get("period_end") or "")
        if not start or not end:
            return []
        values_by_date = {
            item.get("date"): item.get("values") or {}
            for item in statement.get("daily_rows") or []
        }
        rows = []
        cursor = start
        while cursor <= end:
            iso = cursor.isoformat()
            rows.append({"date": iso, "values": values_by_date.get(iso, {})})
            cursor = date.fromordinal(cursor.toordinal() + 1)
        return rows

    month_parts = str(statement.get("month") or "").split("-")
    year = month_parts[0] if month_parts else ""
    month_number = int(month_parts[1]) if len(month_parts) > 1 and month_parts[1].isdigit() else 0
    month_label = f"{year}년 {month_number}월"
    period_text = f"{compact_date(statement.get('period_start'))} 부터 {compact_date(statement.get('period_end'))} 까지"

    summary_lookup = {
        (item["service_code"], item["category_code"]): item
        for item in statement["summary_rows"]
    }
    pdf_statement_order = [
        (SERVICE_W12, CATEGORY_SSG),
        (SERVICE_W12, CATEGORY_COMMON),
        (SERVICE_W12, CATEGORY_TRADERS),
        (SERVICE_W12, CATEGORY_YES24),
        (SERVICE_N3, CATEGORY_SSG),
        (SERVICE_N3, CATEGORY_YES24),
        (SERVICE_W4, CATEGORY_SSG),
        (SERVICE_W4, CATEGORY_COMMON),
        (SERVICE_W4, CATEGORY_YES24),
    ]
    service_first_labels = {
        0: "W1,2_당일",
        4: "N3_당일",
        6: "W4_새벽",
    }

    grid = [["" for _ in range(15)] for _ in range(60)]
    grid[0][0] = month_label
    grid[1][0] = "운송료 지급명세서"
    grid[3][0] = "수수료 명."
    grid[3][1] = "CJ 대한통운 ONE 배송건"
    grid[4][0] = "지급 대상자."
    grid[4][1] = statement["driver"]["name"]
    grid[5][0] = "정산 기간일."
    grid[5][1] = period_text
    grid[6][0] = "지급 예정일."
    grid[6][1] = korean_date(statement.get("payment_due_date"))

    offset_rows = _statement_offset_rows(statement)
    for index, row in enumerate(offset_rows[:24]):
        grid[index][9] = row["waybill"]
        grid[index][10] = row["reason"]
        grid[index][11] = number(row["amount"])
        grid[index][12] = row["product"]
    if len(offset_rows) > 24:
        grid[24][9] = f"외 {len(offset_rows) - 24:,}건"
    if offset_rows:
        grid[25][11] = number(sum(row["amount"] for row in offset_rows))

    grid[8][0] = "지급내역"
    grid[9][1] = "구 분"
    grid[9][3] = "착지 / 건"
    grid[9][5] = "추가박스"
    grid[9][7] = "금액 (부가세 포함)"
    grid[10][0] = "전체 운송료"

    for index, (service_code, category_code) in enumerate(pdf_statement_order):
        row_index = 10 + index
        item = summary_lookup.get((service_code, category_code), {
            "service_code": service_code,
            "category_code": category_code,
            "category_label": CATEGORY_LABELS.get(category_code, category_code),
            "households": 0,
            "extra_boxes": 0,
            "amount": 0,
        })
        grid[row_index][1] = service_first_labels.get(index, "")
        grid[row_index][2] = item["category_label"]
        grid[row_index][3] = number(item["households"])
        grid[row_index][5] = extra_box_value(item)
        grid[row_index][7] = number(item["amount"])

    grid[19][1] = "통 합"
    grid[19][2] = "생수"
    grid[19][3] = "0"
    grid[19][5] = "0"
    grid[19][7] = number(manual_amount("생수"))
    grid[20][1] = "운송료 외"
    grid[20][2] = "추가 전달"
    grid[20][3] = "0"
    grid[20][5] = "-"
    grid[20][7] = number(manual_amount("추가 전달"))
    grid[21][2] = "수당"
    grid[21][3] = "-"
    grid[21][5] = "-"
    grid[21][7] = number(manual_amount("수당"))
    grid[22][0] = "사고귀책 상계건"
    grid[22][6] = "-"
    grid[22][7] = number(manual_amount("사고귀책"))
    grid[24][0] = "지급 합계"
    grid[24][7] = number(statement["total_amount"])

    grid[26][0] = "일 자"
    grid[26][1] = "구 분"
    grid[26][2] = "W1.2_당일"
    grid[26][7] = "구 분"
    grid[26][8] = "N3_당일"
    grid[26][10] = "구 분"
    grid[26][11] = "4W_새벽"
    grid[27][1] = "SSG"
    grid[27][3] = "트레이더스"
    grid[27][5] = "공동배송"
    grid[27][6] = "YES24"
    grid[27][7] = "SSG"
    grid[27][9] = "YES24"
    grid[27][10] = "SSG"
    grid[27][12] = "공동배송"
    grid[27][14] = "YES24"
    headers = [
        "착지/건", "추가박스", "착지/건", "추가박스", "착지/건", "착지 / 건",
        "착지/건", "추가박스", "착지 / 건", "착지/건", "추가박스", "착지/건",
        "추가박스", "착지 / 건",
    ]
    for col_index, value in enumerate(headers, start=1):
        grid[28][col_index] = value

    for row in daily_rows_for_month():
        parsed = parse_date(row["date"])
        if not parsed:
            continue
        grid_row = parsed.day + 28
        if not 29 <= grid_row <= 59:
            continue
        grid[grid_row][0] = short_date(row["date"])
        values = [
            daily_value(row, f"{SERVICE_W12}:{CATEGORY_SSG}", "households"),
            daily_value(row, f"{SERVICE_W12}:{CATEGORY_SSG}", "extra_boxes"),
            daily_value(row, f"{SERVICE_W12}:{CATEGORY_TRADERS}", "households"),
            daily_value(row, f"{SERVICE_W12}:{CATEGORY_TRADERS}", "extra_boxes"),
            daily_value(row, f"{SERVICE_W12}:{CATEGORY_COMMON}", "households"),
            daily_value(row, f"{SERVICE_W12}:{CATEGORY_YES24}", "households"),
            daily_value(row, f"{SERVICE_N3}:{CATEGORY_SSG}", "households"),
            daily_value(row, f"{SERVICE_N3}:{CATEGORY_SSG}", "extra_boxes"),
            daily_value(row, f"{SERVICE_N3}:{CATEGORY_YES24}", "households"),
            daily_value(row, f"{SERVICE_W4}:{CATEGORY_SSG}", "households"),
            daily_value(row, f"{SERVICE_W4}:{CATEGORY_SSG}", "extra_boxes"),
            daily_value(row, f"{SERVICE_W4}:{CATEGORY_COMMON}", "households"),
            daily_value(row, f"{SERVICE_W4}:{CATEGORY_COMMON}", "extra_boxes"),
            daily_value(row, f"{SERVICE_W4}:{CATEGORY_YES24}", "households"),
        ]
        for col_index, value in enumerate(values, start=1):
            grid[grid_row][col_index] = value

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A3),
        rightMargin=5 * mm,
        leftMargin=5 * mm,
        topMargin=5 * mm,
        bottomMargin=5 * mm,
    )
    col_widths = [57] + [69] * 14
    row_heights = [13, 18, 8, 12, 12, 12, 12, 8, 12, 12] + [12] * 14 + [13, 9, 12, 12, 12] + [10] * 31
    sheet = Table(grid, colWidths=col_widths, rowHeights=row_heights, hAlign="LEFT")
    style = [
        ("FONTNAME", (0, 0), (-1, -1), font_name),
        ("FONTSIZE", (0, 0), (-1, -1), 6.3),
        ("LEADING", (0, 0), (-1, -1), 7),
        ("LEFTPADDING", (0, 0), (-1, -1), 2),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, -1), 1),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("FONTSIZE", (9, 0), (14, 25), 5.6),
        ("LEADING", (9, 0), (14, 25), 6),
        ("ALIGN", (9, 0), (10, 25), "LEFT"),
        ("ALIGN", (11, 0), (11, 25), "RIGHT"),
        ("ALIGN", (12, 0), (14, 25), "LEFT"),
        ("FONTNAME", (11, 25), (11, 25), bold_font_name),
        ("FONTNAME", (0, 0), (1, 1), bold_font_name),
        ("FONTSIZE", (0, 0), (1, 0), 11),
        ("FONTSIZE", (0, 1), (1, 1), 16),
        ("FONTNAME", (0, 3), (2, 6), bold_font_name),
        ("FONTSIZE", (0, 3), (2, 6), 8),
        ("GRID", (0, 8), (8, 24), 0.45, colors.black),
        ("GRID", (0, 26), (14, 59), 0.4, colors.black),
        ("BACKGROUND", (0, 8), (8, 8), colors.black),
        ("TEXTCOLOR", (0, 8), (8, 8), colors.white),
        ("BACKGROUND", (0, 23), (8, 23), colors.black),
        ("BACKGROUND", (0, 22), (8, 22), colors.HexColor("#d9d9d9")),
        ("BACKGROUND", (0, 26), (0, 28), colors.black),
        ("TEXTCOLOR", (0, 26), (0, 28), colors.white),
        ("BACKGROUND", (2, 26), (6, 26), colors.black),
        ("TEXTCOLOR", (2, 26), (6, 26), colors.white),
        ("BACKGROUND", (8, 26), (9, 26), colors.black),
        ("TEXTCOLOR", (8, 26), (9, 26), colors.white),
        ("BACKGROUND", (11, 26), (14, 26), colors.black),
        ("TEXTCOLOR", (11, 26), (14, 26), colors.white),
        ("ALIGN", (0, 8), (14, 28), "CENTER"),
        ("ALIGN", (3, 10), (8, 24), "RIGHT"),
        ("ALIGN", (1, 29), (14, 59), "RIGHT"),
        ("FONTNAME", (0, 8), (14, 28), bold_font_name),
        ("FONTNAME", (1, 10), (1, 18), bold_font_name),
        ("FONTNAME", (0, 24), (8, 24), bold_font_name),
        ("SPAN", (0, 0), (1, 0)),
        ("SPAN", (0, 1), (1, 1)),
        ("SPAN", (1, 3), (2, 3)),
        ("SPAN", (1, 4), (2, 4)),
        ("SPAN", (1, 5), (2, 5)),
        ("SPAN", (1, 6), (2, 6)),
        ("SPAN", (0, 8), (8, 8)),
        ("SPAN", (1, 9), (2, 9)),
        ("SPAN", (3, 9), (4, 9)),
        ("SPAN", (5, 9), (6, 9)),
        ("SPAN", (7, 9), (8, 9)),
        ("SPAN", (0, 10), (0, 18)),
        ("SPAN", (1, 10), (1, 13)),
        ("SPAN", (1, 14), (1, 15)),
        ("SPAN", (1, 16), (1, 18)),
        ("SPAN", (3, 10), (4, 10)),
        ("SPAN", (5, 10), (6, 10)),
        ("SPAN", (7, 10), (8, 10)),
        ("SPAN", (3, 11), (4, 11)),
        ("SPAN", (5, 11), (6, 11)),
        ("SPAN", (7, 11), (8, 11)),
        ("SPAN", (3, 12), (4, 12)),
        ("SPAN", (5, 12), (6, 12)),
        ("SPAN", (7, 12), (8, 12)),
        ("SPAN", (3, 13), (4, 13)),
        ("SPAN", (5, 13), (6, 13)),
        ("SPAN", (7, 13), (8, 13)),
        ("SPAN", (3, 14), (4, 14)),
        ("SPAN", (5, 14), (6, 14)),
        ("SPAN", (7, 14), (8, 14)),
        ("SPAN", (3, 15), (4, 15)),
        ("SPAN", (5, 15), (6, 15)),
        ("SPAN", (7, 15), (8, 15)),
        ("SPAN", (3, 16), (4, 16)),
        ("SPAN", (5, 16), (6, 16)),
        ("SPAN", (7, 16), (8, 16)),
        ("SPAN", (3, 17), (4, 17)),
        ("SPAN", (5, 17), (6, 17)),
        ("SPAN", (7, 17), (8, 17)),
        ("SPAN", (3, 18), (4, 18)),
        ("SPAN", (5, 18), (6, 18)),
        ("SPAN", (7, 18), (8, 18)),
        ("SPAN", (3, 19), (4, 19)),
        ("SPAN", (5, 19), (6, 19)),
        ("SPAN", (7, 19), (8, 19)),
        ("SPAN", (1, 20), (1, 21)),
        ("SPAN", (3, 20), (4, 20)),
        ("SPAN", (5, 20), (6, 20)),
        ("SPAN", (7, 20), (8, 20)),
        ("SPAN", (3, 21), (4, 21)),
        ("SPAN", (5, 21), (6, 21)),
        ("SPAN", (7, 21), (8, 21)),
        ("SPAN", (0, 22), (2, 22)),
        ("SPAN", (7, 22), (8, 22)),
        ("SPAN", (0, 23), (8, 23)),
        ("SPAN", (0, 24), (2, 24)),
        ("SPAN", (7, 24), (8, 24)),
        ("SPAN", (0, 26), (0, 28)),
        ("SPAN", (2, 26), (6, 26)),
        ("SPAN", (8, 26), (9, 26)),
        ("SPAN", (11, 26), (14, 26)),
        ("SPAN", (1, 27), (2, 27)),
        ("SPAN", (3, 27), (4, 27)),
        ("SPAN", (7, 27), (8, 27)),
        ("SPAN", (10, 27), (11, 27)),
        ("SPAN", (12, 27), (13, 27)),
    ]
    for row_index in (14, 16):
        style.append(("LINEABOVE", (1, row_index), (8, row_index), 0.8, colors.black))
    for row_index in range(10, 19):
        style.append(("LINEABOVE", (1, row_index), (8, row_index), 0.25, colors.black))
    style.extend([
        ("FONTSIZE", (9, 0), (14, 25), 5.6),
        ("LEADING", (9, 0), (14, 25), 6),
        ("ALIGN", (9, 0), (10, 25), "LEFT"),
        ("ALIGN", (11, 0), (11, 25), "RIGHT"),
        ("ALIGN", (12, 0), (14, 25), "LEFT"),
        ("FONTNAME", (11, 25), (11, 25), bold_font_name),
    ])
    sheet.setStyle(TableStyle(style))
    doc.build([sheet])
    output.seek(0)
    return output.getvalue()
