import shutil
import tempfile
from io import BytesIO
from pathlib import Path

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from openpyxl import Workbook

from .constants import CATEGORY_COMMON, CATEGORY_SSG, CATEGORY_TRADERS, SERVICE_N3, SERVICE_W12, SERVICE_W4
from .models import OneDriver, OneDriverStatementOverride, OneSettlementUpload, OneShipmentOrder
from .services import build_month_summary, delete_one_upload, import_one_file, list_drivers, parse_one_workbook, parse_one_workbook_bundle


TEST_MEDIA_ROOT = tempfile.mkdtemp(prefix="one-settlement-test-media-")


def make_workbook_bytes(rows):
    workbook = Workbook()
    raw = workbook.active
    raw.title = "RAW"
    raw.append(["일자", "주문번호", "운송장번호", "수수료명칭", "시군구", "SM"])
    raw.append([None, None, None, "명", None, None])
    for row in rows:
        raw.append(row)

    ignored = workbook.create_sheet("생수RAW")
    ignored.append(["일자", "주문번호", "수수료명칭", "시군구", "SM"])
    ignored.append(["2026-05-01", "IGNORED", "새벽_SSG_N3 에프엔로지스", "종로구", "무시대상"])

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def make_workbook_with_water_bytes(raw_driver_name="홍길동", water_driver_name="홍길동"):
    workbook = Workbook()
    raw = workbook.active
    raw.title = "RAW"
    raw.append(["일자", "주문번호", "운송장번호", "수수료명칭", "시군구", "SM"])
    raw.append([None, None, None, "명", None, None])
    raw.append(["2026-05-01", "ORDER-1", "WAY-1", "당일_SSG N3 3W_에프엔로지스", "마포구", raw_driver_name])

    water = workbook.create_sheet("생수RAW")
    water.append(["배송협력사", "배송일자", "SM", "운송장번호", "주문번호", "받는분주소", "품목명"])
    water.append(["당일)N3_에프엔로지스", "2026-05-01", water_driver_name, "WATER-1", "ORDER-W1", "서울", "삼다수 2L"])
    water.append(["당일)N3_에프엔로지스", "2026-05-01", water_driver_name, "WATER-2", "ORDER-W2", "서울", "삼다수 500ml"])

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def make_april_style_workbook_bytes():
    workbook = Workbook()
    raw = workbook.active
    raw.title = "로우"
    raw.append(["일자", "주문번호", "운송장번호", None, "수수료명칭", "시군구", "건별수수료", None, None, None, None, None, "최종처리시간", "SM"])
    raw.append([None, None, None, "명", None, None, "기준단가", "할증유형1(물량)", "할증유형2(물성)", "할증유형3(기타)", "합계", "명", None, None])
    raw.append(["2026-04-01", "ORDER-1", "WAY-1", "당일)N3_에프엔로지스", "당일_SSG N3 2W_에프엔로지스", "마포구", "1450", "0", "0", "0", "1450", "SSG_N3센터", "2026-04-01 20:56:18", "홍길동"])
    raw.append(["2026-04-01", "ORDER-1", "WAY-2", "당일)N3_에프엔로지스", "당일_SSG N3 2W_에프엔로지스", "마포구", "1450", "0", "0", "0", "1450", "SSG_N3센터", "2026-04-01 20:56:18", "홍길동"])

    water = workbook.create_sheet("생수")
    water.append([None] * 14)
    water.append([None, "구분", "당일)N3_에프엔로지스", None, None, "배송협력사", "배송일자", "SM", "운송장번호", "주문번호", "받는분주소", "품목명", "주관고객명", "발송고객명"])
    water.append([None, None, "생수 수량", "생수 수수료", None, "당일)N3_에프엔로지스", "2026-04-01", "홍길동", "WATER-1", "ORDER-W1", "서울", "삼다수 2L", "SSG", "SSG"])

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


def make_position_based_workbook_bytes():
    workbook = Workbook()
    raw = workbook.active
    raw.title = "5월 원본"
    raw.append(["배송 일자", "주문 번호", "송장 번호", "수수료 명칭", "시/군/구", "배송원명"])
    raw.append(["2026-05-02", "ORDER-P1", "WAY-P1", "새벽_SSG_N3 에프엔로지스", "종로구", "홍길동"])

    water = workbook.create_sheet("두번째시트")
    water.append(["설명 행"])
    water.append(["배송 협력사", "배송 일자", "배송원명", "송장 번호", "주문 번호", "배송 주소", "상품명"])
    water.append(["새벽)종로_에프엔로지스", "2026-05-02", "홍길동", "WATER-P1", "ORDER-WP1", "서울", "생수"])

    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


class OneSettlementParserTests(TestCase):
    def test_parse_raw_only_groups_household_boxes_without_jongno_extra(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "새벽_SSG_N3 에프엔로지스", "종로구", "홍길동"],
            ["2026-05-01", "ORDER-1", "WAY-2", "새벽_SSG_N3 에프엔로지스", "종로구", "홍길동"],
            ["2026-05-01", "ORDER-2", "WAY-3", "새벽_SSG_N3 에프엔로지스", "강남구", "홍길동"],
        ])

        _raw_hash, total_rows, orders = parse_one_workbook(content)

        self.assertEqual(total_rows, 3)
        self.assertEqual(len(orders), 2)
        first = next(order for order in orders if order.order_number == "ORDER-1")
        self.assertEqual(first.boxes, 2)
        self.assertEqual(first.extra_boxes, 1)
        self.assertEqual(first.city, "종로구")
        self.assertEqual(first.service_code, SERVICE_W4)
        self.assertEqual(first.category_code, CATEGORY_SSG)
        self.assertEqual(int(first.base_amount), 2800)
        self.assertEqual(int(first.jongno_extra_amount), 0)
        self.assertEqual(int(first.amount), 2800)

    def test_w12_common_and_traders_use_daily_volume_rate(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "당일_공동배송망 에프엔로지스", "마포구", "홍길동"],
            ["2026-05-01", "ORDER-1", "WAY-2", "당일_공동배송망 에프엔로지스", "마포구", "홍길동"],
            ["2026-05-01", "ORDER-2", "WAY-3", "당일_트레이더스 에프엔로지스", "마포구", "홍길동"],
            ["2026-05-01", "ORDER-2", "WAY-4", "당일_트레이더스 에프엔로지스", "마포구", "홍길동"],
            ["2026-05-01", "ORDER-2", "WAY-5", "당일_트레이더스 에프엔로지스", "마포구", "홍길동"],
        ])

        _raw_hash, _total_rows, orders = parse_one_workbook(content)
        common = next(order for order in orders if order.order_number == "ORDER-1")
        traders = next(order for order in orders if order.order_number == "ORDER-2")

        self.assertEqual(common.service_code, SERVICE_W12)
        self.assertEqual(common.category_code, CATEGORY_COMMON)
        self.assertEqual(int(common.amount), 3300)
        self.assertEqual(traders.category_code, CATEGORY_TRADERS)
        self.assertEqual(int(traders.amount), 3650)

    def test_parse_blank_order_number_uses_recent_matching_waybill_group(self):
        content = make_workbook_bytes([
            ["2026-05-01", None, "WAY-1", "새벽_공동배송망 에프엔로지스", "마포구", "홍길동"],
            ["2026-05-01", None, "WAY-2", "새벽_공동배송망 에프엔로지스", "마포구", "홍길동"],
        ])

        _raw_hash, total_rows, orders = parse_one_workbook(content)

        self.assertEqual(total_rows, 2)
        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0].order_number, "WAY-1")
        self.assertEqual(orders[0].boxes, 2)
        self.assertEqual(orders[0].extra_boxes, 1)
        self.assertEqual(int(orders[0].amount), 2800)

    def test_n3_fee_name_takes_precedence_over_2w_text(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "당일_SSG N3 2W_에프엔로지스", "마포구", "홍길동"],
        ])

        _raw_hash, _total_rows, orders = parse_one_workbook(content)

        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0].service_code, SERVICE_N3)
        self.assertEqual(orders[0].category_code, CATEGORY_SSG)

    def test_raw_driver_name_is_not_normalized(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "당일_SSG N3 3W_에프엔로지스", "마포구", "이현국A"],
            ["2026-05-01", "ORDER-2", "WAY-2", "당일_SSG N3 3W_에프엔로지스", "마포구", "최길준_A"],
        ])

        _raw_hash, _total_rows, orders = parse_one_workbook(content)

        self.assertEqual({order.driver_name for order in orders}, {"이현국A", "최길준_A"})

    def test_raw_driver_name_alias_is_not_applied(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "당일_SSG N3 3W_에프엔로지스", "마포구", "\uc815\uc2b9\ud638"],
            ["2026-05-01", "ORDER-2", "WAY-2", "당일_SSG N3 3W_에프엔로지스", "마포구", "\ucd5c\uae30\uc900"],
        ])

        _raw_hash, _total_rows, orders = parse_one_workbook(content)

        self.assertEqual({order.driver_name for order in orders}, {"\uc815\uc2b9\ud638", "\ucd5c\uae30\uc900"})

    def test_parse_water_sheet_as_auto_manual_items(self):
        content = make_workbook_with_water_bytes()

        _raw_hash, _total_rows, _orders, manual_items = parse_one_workbook_bundle(content)

        self.assertEqual(len(manual_items), 2)
        self.assertEqual({item.driver_name for item in manual_items}, {"홍길동"})
        self.assertEqual(sum(int(item.amount) for item in manual_items), 400)

    def test_water_driver_alias_matches_existing_raw_driver_only(self):
        content = make_workbook_with_water_bytes(raw_driver_name="\uc815\uc2b9\uc6a9", water_driver_name="\uc815\uc2b9\ud638")

        _raw_hash, _total_rows, orders, manual_items = parse_one_workbook_bundle(content)

        self.assertEqual({order.driver_name for order in orders}, {"\uc815\uc2b9\uc6a9"})
        self.assertEqual({item.driver_name for item in manual_items}, {"\uc815\uc2b9\uc6a9"})

    def test_april_one_workbook_aliases_parse_raw_and_water_sheets(self):
        content = make_april_style_workbook_bytes()

        _raw_hash, total_rows, orders, manual_items = parse_one_workbook_bundle(content)

        self.assertEqual(total_rows, 2)
        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0].service_code, SERVICE_N3)
        self.assertEqual(orders[0].category_code, CATEGORY_SSG)
        self.assertEqual(orders[0].boxes, 2)
        self.assertEqual(len(manual_items), 1)
        self.assertEqual(manual_items[0].driver_name, "홍길동")

    def test_position_based_sheet_fallback_and_header_aliases(self):
        content = make_position_based_workbook_bytes()

        _raw_hash, total_rows, orders, manual_items = parse_one_workbook_bundle(content)

        self.assertEqual(total_rows, 1)
        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0].order_number, "ORDER-P1")
        self.assertEqual(orders[0].driver_name, "홍길동")
        self.assertEqual(len(manual_items), 1)
        self.assertEqual(manual_items[0].payload["waybill"], "WATER-P1")


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class OneSettlementImportTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)

    def test_duplicate_raw_hash_is_not_imported_twice(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "새벽_SSG_N3 에프엔로지스", "종로구", "홍길동"],
        ])

        first = import_one_file(
            company_app="new",
            uploaded_by=None,
            uploaded_file=SimpleUploadedFile("one-1.xlsx", content),
        )
        second = import_one_file(
            company_app="new",
            uploaded_by=None,
            uploaded_file=SimpleUploadedFile("one-copy.xlsx", content),
        )

        self.assertFalse(first["duplicate"])
        self.assertTrue(second["duplicate"])
        self.assertEqual(first["upload"].id, second["upload"].id)
        self.assertEqual(OneSettlementUpload.objects.count(), 1)
        self.assertEqual(OneShipmentOrder.objects.count(), 1)
        self.assertEqual(OneDriver.objects.filter(company_app="new", name="홍길동").count(), 1)

    def test_import_recalculates_w12_daily_rate_across_month_uploads(self):
        first_rows = [
            ["2026-05-01", f"ORDER-{index:03d}", f"WAY-{index:03d}", "당일_공동배송망 에프엔로지스", "마포구", "홍길동"]
            for index in range(1, 500)
        ]
        second_rows = [
            ["2026-05-01", "ORDER-500", "WAY-500", "당일_공동배송망 에프엔로지스", "마포구", "홍길동"]
        ]

        import_one_file(
            company_app="new",
            uploaded_by=None,
            uploaded_file=SimpleUploadedFile("one-499.xlsx", make_workbook_bytes(first_rows)),
        )
        self.assertEqual(
            sum(int(order.amount) for order in OneShipmentOrder.objects.filter(category_code=CATEGORY_COMMON)),
            499 * 1650,
        )

        import_one_file(
            company_app="new",
            uploaded_by=None,
            uploaded_file=SimpleUploadedFile("one-500.xlsx", make_workbook_bytes(second_rows)),
        )

        self.assertEqual(OneShipmentOrder.objects.filter(category_code=CATEGORY_COMMON).count(), 500)
        self.assertEqual(
            sum(int(order.amount) for order in OneShipmentOrder.objects.filter(category_code=CATEGORY_COMMON)),
            500 * 1630,
        )

    def test_delete_upload_removes_orders_and_auto_water_manual_items(self):
        content = make_workbook_with_water_bytes()
        result = import_one_file(
            company_app="new",
            uploaded_by=None,
            uploaded_file=SimpleUploadedFile("one-water.xlsx", content),
        )
        driver = OneDriver.objects.get(company_app="new", name="홍길동")
        override = OneDriverStatementOverride.objects.get(company_app="new", month="2026-05", driver=driver)

        self.assertEqual(OneSettlementUpload.objects.count(), 1)
        self.assertEqual(OneShipmentOrder.objects.count(), 1)
        self.assertEqual(sum(int(item["amount"]) for item in override.manual_items), 400)

        delete_one_upload(company_app="new", upload_id=result["upload"].id)
        override.refresh_from_db()

        self.assertEqual(OneSettlementUpload.objects.count(), 0)
        self.assertEqual(OneShipmentOrder.objects.count(), 0)
        self.assertEqual(override.manual_items, [])

    def test_summary_uses_statement_total_without_manual_only_drivers(self):
        content = make_workbook_bytes([
            ["2026-05-01", "ORDER-1", "WAY-1", "당일_SSG N3 3W_에프엔로지스", "마포구", "홍길동"],
        ])
        import_one_file(
            company_app="new",
            uploaded_by=None,
            uploaded_file=SimpleUploadedFile("one.xlsx", content),
        )

        active_driver = OneDriver.objects.get(company_app="new", name="홍길동")
        manual_only_driver = OneDriver.objects.create(company_app="new", name="수기전용")
        OneDriverStatementOverride.objects.create(
            company_app="new",
            shipper_code="one",
            month="2026-05",
            driver=active_driver,
            manual_items=[{"label": "추가 전달", "amount": 1000}],
        )
        OneDriverStatementOverride.objects.create(
            company_app="new",
            shipper_code="one",
            month="2026-05",
            driver=manual_only_driver,
            manual_items=[{"label": "지원금", "amount": 9999}],
        )

        summary = build_month_summary("new", "2026-05")["summary"]
        drivers = list_drivers("new", "2026-05")

        self.assertEqual(summary["driver_count"], 1)
        self.assertEqual(summary["manual_total"], 1000)
        self.assertEqual(summary["statement_total_amount"], summary["total_amount"] + 1000)
        self.assertEqual([driver["name"] for driver in drivers], ["홍길동"])
        self.assertEqual(drivers[0]["manual_total"], 1000)
        self.assertEqual(drivers[0]["statement_total_amount"], drivers[0]["amount"] + 1000)


class OneSettlementSampleDataTests(TestCase):
    def test_imjunhyeong_may_2026_w4_ssg_matches_reference_sample(self):
        data_root = Path(__file__).resolve().parents[3] / "ONE" / "data"
        if not data_root.exists():
            self.skipTest("ONE/data sample files are not available")

        households = 0
        extra_boxes = 0
        amount = 0
        seen_hashes = set()

        for path in sorted(data_root.rglob("*.xlsx")):
            if "(새벽)" not in str(path):
                continue
            raw_hash, _total_rows, orders = parse_one_workbook(path.read_bytes())
            if raw_hash in seen_hashes:
                continue
            seen_hashes.add(raw_hash)
            for order in orders:
                if (
                    order.driver_name == "임준형"
                    and order.service_code == SERVICE_W4
                    and order.category_code == CATEGORY_SSG
                ):
                    households += 1
                    extra_boxes += order.extra_boxes
                    amount += int(order.amount)

        self.assertGreater(len(seen_hashes), 0)
        self.assertEqual(households, 993)
        self.assertEqual(extra_boxes, 1008)
        self.assertEqual(amount, 2783400)
